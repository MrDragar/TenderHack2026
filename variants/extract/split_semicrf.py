"""Слагаемые формулы split_probability с обучаемыми весами: полумарковская CRF на решётке фраз.

Та же решётка, что в split_probability: каждая фраза между запятыми либо просто фраза, либо
кусок из нескольких фраз «название + значение» с разрезом по пробелу. Счёт куска это сумма
десяти признаков с весами: log P(название), -log P(фон) у слов названия, log P(значение |
название), -log P(фон) у слов значения, запятые в значении и в названии, свободный член,
длина значения в словах, есть ли в значении число или единица, известно ли название целиком.
Веса стартуют с ручной формулы и учатся максимумом условного правдоподобия по прямому и
обратному проходу (Sarawagi, Cohen 2004) на склейках листа 1 вне замера. Ответ это лучший
разбор, пару отдаём, если её маргинал не меньше 1/2.
Предсказание: меньше «значение забирает следующую пару» и недобранных списков, верные
не теряются.
"""

import json
import math
import zlib
from functools import cache
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize

from contracts.pipeline import Fragment
from pipeline.run import portal
from pipeline.text import PAIR, words
from scripts.products import RAW_DIR, SAMPLE, SHEETS, section_names, split_pairs
from variants.extract.split_probability import (
    CLAUSE,
    NUMBER,
    SPAN,
    fragment,
    model,
    numbers,
    pair,
    prob,
    split,
    stripped,
    tokenize,
)

WEIGHTS_PATH = Path("data/processed/extract_semicrf.json")
# Те же склейки, что учат crf: файл общий, кто первый, тот и собрал.
TRAIN_PATH = Path("data/processed/extract_crf_train.json")
# Каждый двухсотый товар вне замера: около полутора тысяч, обучение укладывается в минуты.
EVERY = 200
# Неверная пара стоит вдвое дороже пропущенной: отвечаем, если уверены хотя бы наполовину.
MIN_PROBA = 0.5
FEATURES = 10


def extract(text: str) -> list[Fragment]:
    found = []
    pairs = []
    for match in PAIR.finditer(text):
        found.append(pair(text, match))
        pairs.append((match.start("name"), match.end("value")))
    for run in runs(text, pairs):
        for start, end in run:
            found.extend(numbers(text, start, end))
            if words(text[start:end]):
                found.append(fragment("words", None, text, start, end))
        found.extend(segment(text, run))
    found.sort(key=lambda f: f["start"])
    return found


def runs(text: str, pairs: list[tuple[int, int]]) -> list[list[tuple[int, int]]]:
    found: list[list[tuple[int, int]]] = [[]]
    for clause in CLAUSE.finditer(text):
        start, end = stripped(text, clause.start(), clause.end())
        if start == end:
            continue
        if any(s <= start < e for s, e in pairs):
            found.append([])
            continue
        found[-1].append((start, end))
    return [run for run in found if run]


def segment(text: str, clauses: list[tuple[int, int]]) -> list[Fragment]:
    segs, feats = lattice(text, clauses)
    if not len(segs):
        return []
    scores = feats @ weights()
    starts, first, last = groups(segs)
    span = group_logsumexp(scores, starts)
    alpha, beta = passes(len(clauses), first, last, span)
    found = []
    for g in viterbi(len(clauses), first, last, span):
        lo = starts[g]
        hi = starts[g + 1] if g + 1 < len(starts) else len(segs)
        k = lo + int(np.argmax(scores[lo:hi]))
        sure = math.exp(alpha[first[g]] + scores[k] + beta[last[g] + 1] - alpha[-1])
        if sure >= MIN_PROBA:
            found.append(split(text, clauses[first[g]][0], int(segs[k, 2]), clauses[last[g]][1]))
    return found


def lattice(text: str, clauses: list[tuple[int, int]]) -> tuple[np.ndarray, np.ndarray]:
    """Все куски (первая фраза, последняя фраза, разрез) не длиннее SPAN фраз, по порядку
    последней фразы, затем первой: разрезы одного отрезка идут подряд."""
    segs = []
    feats = []
    names: dict[str, tuple[float, float, float]] = {}
    tokens: dict[tuple[str, str], tuple[float, float]] = {}
    for j in range(len(clauses)):
        for i in range(max(0, j + 1 - SPAN), j + 1):
            start, end = clauses[i][0], clauses[j][1]
            k = i
            for cut in range(start + 1, end):
                if k < j and cut >= clauses[k + 1][0]:
                    k += 1
                if text[cut - 1].isspace() and not text[cut].isspace():
                    segs.append((i, j, cut))
                    name = text[start:cut].strip()
                    if name not in names:
                        names[name] = name_part(name)
                    item = features(name, names[name], text[cut:end], tokens)
                    feats.append([*item, j - k, k - i])
    return np.array(segs, dtype=np.int32).reshape(-1, 3), np.array(feats).reshape(-1, FEATURES)


def features(
    name: str,
    name_scores: tuple[float, float, float],
    value: str,
    seen: dict[tuple[str, str], tuple[float, float]],
) -> list[float]:
    key = name.casefold()
    value_like = 0.0
    value_bg = 0.0
    tokens = tokenize(value)
    for token in tokens:
        if (token, key) not in seen:
            seen[token, key] = (value_word(token, key), math.log(prob(token, "background")))
        like, bg = seen[token, key]
        value_like += like
        value_bg += bg
    name_like, name_bg, known = name_scores
    return [
        name_like,
        -name_bg,
        value_like,
        -value_bg,
        1.0,
        len(tokens),
        float(NUMBER.search(value) is not None),
        known,
    ]


def name_part(name: str) -> tuple[float, float, float]:
    stats = model()
    tokens = tokenize(name)
    count = stats["names"].get(name.casefold(), 0)
    if count:
        likely = math.log(count / stats["names_total"])
    else:
        likely = math.log(stats["new_name"]) + sum(math.log(prob(t, "name_words")) for t in tokens)
    return likely, sum(math.log(prob(t, "background")) for t in tokens), float(count > 0)


def value_word(token: str, name: str) -> float:
    own = model()["values"].get(name)
    likely = prob(token, "value_words")
    if own:
        likely = (own[token] + len(own) * likely) / (own.total() + len(own))
    return math.log(likely)


def groups(segs: np.ndarray) -> tuple[np.ndarray, list[int], list[int]]:
    """Начала отрезков (первая фраза, последняя фраза) в списке кусков и сами отрезки."""
    edge = np.flatnonzero(np.any(np.diff(segs[:, :2], axis=0) != 0, axis=1)) + 1
    starts = np.concatenate([[0], edge])
    return starts, segs[starts, 0].tolist(), segs[starts, 1].tolist()


def group_logsumexp(scores: np.ndarray, starts: np.ndarray) -> np.ndarray:
    top = np.maximum.reduceat(scores, starts)
    sizes = np.diff(np.append(starts, len(scores)))
    return top + np.log(np.add.reduceat(np.exp(scores - np.repeat(top, sizes)), starts))


def passes(
    n: int, first: list[int], last: list[int], span: np.ndarray
) -> tuple[list[float], list[float]]:
    """Просто фраза стоит 0, отрезок сумму по своим разрезам; alpha[n] это log Z."""
    span = span.tolist()
    alpha = [0.0] * (n + 1)
    by_start: list[list[int]] = [[] for _ in range(n)]
    g = 0
    for j in range(n):
        alpha[j + 1] = alpha[j]
        while g < len(first) and last[g] == j:
            alpha[j + 1] = log_add(alpha[j + 1], alpha[first[g]] + span[g])
            by_start[first[g]].append(g)
            g += 1
    beta = [0.0] * (n + 1)
    for i in range(n - 1, -1, -1):
        beta[i] = beta[i + 1]
        for g in by_start[i]:
            beta[i] = log_add(beta[i], beta[last[g] + 1] + span[g])
    return alpha, beta


def log_add(a: float, b: float) -> float:
    if a < b:
        a, b = b, a
    return a + math.log1p(math.exp(b - a))


def viterbi(n: int, first: list[int], last: list[int], span: np.ndarray) -> list[int]:
    """Лучший разбор по отрезкам; счёт отрезка тут log-сумма по разрезам, ответ его лучший
    разрез. Возвращает номера отрезков."""
    best = [0.0] * (n + 1)
    back: list[int | None] = [None] * (n + 1)
    g = 0
    for j in range(n):
        best[j + 1] = best[j]
        while g < len(first) and last[g] == j:
            score = best[first[g]] + span[g]
            if score > best[j + 1]:
                best[j + 1], back[j + 1] = score, g
            g += 1
    chosen = []
    j = n
    while j > 0:
        g = back[j]
        if g is None:
            j -= 1
        else:
            chosen.append(g)
            j = first[g]
    return chosen[::-1]


def start_weights() -> np.ndarray:
    """Ручная формула split_probability: с неё начинаем и к ней тянет регуляризация."""
    comma = model()["comma"]
    return np.array([1, 1, 1, 1, 0, 0, 0, 0, comma, comma], dtype=float)


@cache
def weights() -> np.ndarray:
    if not WEIGHTS_PATH.exists():
        train()
    return np.array(json.loads(WEIGHTS_PATH.read_text()))


def train() -> None:
    data = training_runs()
    feats = np.concatenate([d[1] for d in data])
    offsets = np.cumsum([0] + [len(d[1]) for d in data])
    gold = np.zeros(FEATURES)
    for _, run_feats, _, _, _, chosen in data:
        gold += run_feats[chosen].sum(axis=0)
    starts = np.concatenate([d[2] + lo for d, lo in zip(data, offsets, strict=False)])
    sizes = np.diff(np.append(starts, len(feats)))
    bounds = np.cumsum([0] + [len(d[2]) for d in data])
    origin = start_weights()

    def loss(w: np.ndarray) -> tuple[float, np.ndarray]:
        scores = feats @ w
        span = group_logsumexp(scores, starts)
        mass = np.zeros(len(starts))
        log_z = 0.0
        for (n, _, _, first, last, _), lo, hi in zip(data, bounds, bounds[1:], strict=False):
            alpha, beta = passes(n, first, last, span[lo:hi])
            log_z += alpha[n]
            for g in range(hi - lo):
                mass[lo + g] = alpha[first[g]] + beta[last[g] + 1] - alpha[n]
        marginal = np.exp(np.repeat(mass, sizes) + scores)
        shift = w - origin
        return log_z - gold @ w + 0.5 * shift @ shift, marginal @ feats - gold + shift

    result = minimize(loss, origin, jac=True, method="L-BFGS-B", options={"maxiter": 100})
    WEIGHTS_PATH.write_text(json.dumps(result.x.tolist()))


def training_runs() -> list[tuple]:
    """Решётки листа 1 вне замера с номерами эталонных кусков. Прогон, где эталонная пара не
    ложится на решётку (двоеточие, точка, длиннее SPAN фраз), выкидываем целиком."""
    found = []
    for text, gold_pairs in examples():
        colon = [(m.start("name"), m.end("value")) for m in PAIR.finditer(text)]
        for run in runs(text, colon):
            segs, feats = lattice(text, run)
            if not len(segs):
                continue
            index = {}
            for k, (i, j, cut) in enumerate(segs.tolist()):
                index[run[i][0], cut, run[j][1]] = k
            chosen = []
            for name_start, cut, end in gold_pairs:
                if run[0][0] <= name_start < run[-1][1] or run[0][0] < end <= run[-1][1]:
                    chosen.append(index.get((name_start, cut, end)))
            if None not in chosen:
                starts, first, last = groups(segs)
                found.append((len(run), feats, starts, first, last, chosen))
    return found


def examples() -> list[tuple[str, list[tuple[int, int, int]]]]:
    """Склеено как у судьи: название товара, затем «Название значение» через запятую."""
    if not TRAIN_PATH.exists():
        TRAIN_PATH.write_text(json.dumps(read_sheets(), ensure_ascii=False))
    return json.loads(TRAIN_PATH.read_text())[:: EVERY // 20]


def read_sheets() -> list[tuple[str, list[tuple[int, int, int]]]]:
    """Каждый двадцатый товар, не попавший в выборку замера, как в crf."""
    rows = []
    for sheet_name in SHEETS:
        book = load_workbook(RAW_DIR / sheet_name, read_only=True)
        goods, chars = book.worksheets
        sections = section_names(chars.iter_rows(min_row=2, values_only=True), portal())
        names = set().union(*sections.values())
        for row in goods.iter_rows(min_row=2, values_only=True):
            key = zlib.crc32(str(row[0]).encode())
            if not row[0] or not row[1] or not key % SAMPLE or key // SAMPLE % 20:
                continue
            text = " ".join(str(row[1]).split())
            pairs = []
            for name, value in split_pairs(str(row[5] or ""), names):
                text += ", "
                start = len(text)
                pairs.append((start, start + len(name) + 1, start + len(name) + 1 + len(value)))
                text += f"{name} {value}"
            if pairs:
                rows.append((text, pairs))
        book.close()
    return rows
