"""Текст объясняется не жадно, а набором значений с наибольшим суммарным выигрышем над шумом.

Вывод: на dev против noisy_channel (тот же день, тот же генератор) верно 552 → 553, отказ
85 → 85, неверно 46 → 45, цена 0.259 → 0.256; по ключам лучше 1, хуже 0. Исправилось
«искусственный утеплитель и натуральный утеплитель»: жадность отдавала оба «утеплителя»
одному значению. Выигрыш мал, потому что почти в каждом куске одно значение, и там жадный
выбор уже оптимален. Попутно найдено: ничью «основного» против «дополнительного» жадный
выигрывал случайно (P ровно 1/2), честно её решают повторы слов в названии. Время то же,
около 23 мс на пример. Остальные ошибки про числа, «без X» и пометку формы слова.

Зашумлённый канал как в noisy_channel (Brill и Moore 2000), меняется только выбор. Как
в record linkage один-к-одному (Jaro, Sadinle) и глобальном выводе Roth и Yih, ищем максимум
суммы log-выигрышей при ограничениях: слово текста объясняется не больше одного раза (и
словом значения, и словом названия), у характеристики single одно значение, у multiple
несколько, пустой набор это исход «текст ни при чём». Кусков мало, поэтому точный перебор
наборов вместо ILP. Ответ даём по маргиналу: доля веса exp(выигрыш) наборов с этим значением
не меньше 1/2, правило Чоу при ценах 0, 1, 2.
"""

import json
from collections import Counter
from collections.abc import Callable
from difflib import SequenceMatcher
from functools import cache
from math import exp, inf, log
from pathlib import Path
from typing import Literal, TypedDict

from openpyxl import load_workbook

from contracts.pipeline import Candidate, FixKind, FoundChar, Fragment, Rule
from contracts.ref import Characteristic, Portal, PortalChar, Position, Unit, Value
from pipeline.explain import mark
from pipeline.text import Word, stem, stems, words
from pipeline.values import UNIT_ALIASES, bounds, convert, same, scale, show, tidy
from scripts.products import RAW_DIR, SHEETS, products

# Перепутанная единица чаще всего сдвигает число на порядок: см вместо мм, м вместо см.
SHIFTS = (10, 0.1, 100, 0.01, 1000, 0.001)
SHOWN_OPTIONS = 6
COSTS_PATH = Path("data/processed/link_noisy_channel_costs.json")
# Написание базовой характеристики сравниваем с эталонным, только если это одно значение.
SAME_VALUE = 0.5
MAX_EDITS = 3
RUSSIAN = "абвгдежзийклмнопрстуфхцчшщъыьэюя"
ALPHABET = len(RUSSIAN) + 1
OPS = ("sub", "del", "ins", "swap")
# Перебор наборов растёт как 2^n, а в куске текста редко больше нескольких значений.
MAX_OPTIONS = 12

type Source = Literal["ktru", "portal"]


class Option(TypedDict):
    char: Characteristic
    value: str
    source: Source


class Explained(TypedDict):
    option: Option
    gain: float
    used: set[int]
    named: set[int]
    kind: FixKind | None


def link(
    fragments: list[Fragment],
    candidates: list[Candidate],
    section: str | None,
    ref: list[Position],
    portal: Portal,
) -> list[FoundChar]:
    by_code = {p["code"]: p for p in ref}
    positions = [by_code[candidate["code"]] for candidate in candidates]
    chars = [char for position in positions for char in position["characteristics"]]
    allowed = allowed_values(chars)
    portal_chars = [as_ktru(char) for char in portal.get(section or "", [])]
    product = stems(" ".join(c["name"] for c in candidates))
    prior = value_prior(positions)

    pairs = told_pairs(fragments, chars, portal_chars, allowed, product, prior)
    shared = set.intersection(*(stems(c["name"]) for c in candidates)) if candidates else set()
    values = told_values(fragments, chars, portal_chars, product, shared, prior)
    numbers = number_values(fragments, chars, portal_chars, allowed, product)
    found = choose(pairs, values, numbers)
    # Обязательность у соседних кодов разная: обязательна, если обязательна хоть у одного.
    required = {char["name"] for char in chars if char["required"]}
    for char in found:
        char["required"] = char["source"] == "ktru" and char["name"] in required
    return found


def choose(
    pairs: list[FoundChar], values: list[FoundChar], numbers: list[FoundChar]
) -> list[FoundChar]:
    """Явная пара сильнее значения без названия, оно сильнее числа. Значения одного правила
    уже выбраны каналом, между правилами разные значения дают правку «два значения»."""
    by_name: dict[str, list[FoundChar]] = {}
    for rule in (pairs, values, numbers):
        mine: dict[str, list[FoundChar]] = {}
        for char in rule:
            mine.setdefault(char["name"], []).append(char)
        for name, chars in mine.items():
            kept = by_name.get(name)
            if kept is None:
                by_name[name] = chars
                continue
            other = [c for c in chars if not any(same(k["value"], c["value"]) for k in kept)]
            if other:
                mark(kept[0], "conflict", other=other[0]["fragment"])
    return [char for chars in by_name.values() for char in chars]


@cache
def costs() -> dict[str, float]:
    """−log частоты каждой порчи. Счётчики собираются один раз по листам 2 в data/processed."""
    if not COSTS_PATH.exists():
        COSTS_PATH.write_text(json.dumps(mined_counts(), ensure_ascii=False, indent=1))
    counts = json.loads(COSTS_PATH.read_text())
    return priced(Counter(counts["words"]), Counter(counts["edits"]))


@cache
def noise() -> float:
    """Слово шума это порча «лишнее слово» и любое слово из тех, что поставщики пишут в товарах
    листа 1, на выбор."""
    vocabulary = set()
    for product in products({}, build=False):
        vocabulary |= {word["text"] for word in words(product["name"])}
        for name, value in product["pairs"]:
            vocabulary |= {word["text"] for word in words(f"{name} {value}")}
    return costs()["noise"] + log(len(vocabulary))


def mined_counts() -> dict[str, Counter[str]]:
    written: Counter[tuple[str, tuple[str, ...]]] = Counter()
    for sheet in SHEETS:
        book = load_workbook(RAW_DIR / sheet, read_only=True)
        for row in book.worksheets[1].iter_rows(min_row=2, max_col=11, values_only=True):
            linked = tuple(linked_values(str(row[10] or "")))
            if row[1] and row[6] and linked:
                for value in split_values(str(row[6])):
                    written[value, linked] += 1
        book.close()
    words_seen: Counter[str] = Counter()
    edits: Counter[str] = Counter()
    for (value, linked), times in written.items():
        count_spoils(value, linked, times, words_seen, edits)
    return {"words": words_seen, "edits": edits}


def priced(words_seen: Counter[str], edits: Counter[str]) -> dict[str, float]:
    """Цена вида порчи по словам, цена правки по виду и по букве: P(буква текста | буква КТРУ)
    для замены и P(буква) для вставки, со сглаживанием на букву алфавита."""
    said = sum(words_seen[k] for k in ("clean", "layout", "form", "edit", "noise"))
    found = {}
    for kind in ("clean", "layout", "form", "edit", "noise"):
        found[kind] = -log((words_seen[kind] + 1) / said)
    found["omit"] = -log((words_seen["omit"] + 1) / (words_seen["kept"] + words_seen["omit"]))
    by_kind: Counter[str] = Counter()
    by_letter: Counter[str] = Counter()
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        by_kind[kind] += count
        if kind == "sub":
            by_letter["sub:" + letters[0]] += count
    total = sum(by_kind[kind] for kind in OPS) + len(OPS)
    for kind in OPS:
        found[kind] = -log((by_kind[kind] + 1) / total)
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        if kind == "sub":
            found[op] = -log((count + 1) / (by_letter["sub:" + letters[0]] + ALPHABET))
        elif kind == "ins":
            found[op] = -log((count + 1) / (by_kind["ins"] + ALPHABET))
        elif kind == "form":
            found[op] = 0.0
    for letter in RUSSIAN:
        found.setdefault("sub:" + letter, -log(1 / (by_letter["sub:" + letter] + ALPHABET)))
    found["ins:"] = -log(1 / (by_kind["ins"] + ALPHABET))
    return found


def op_price(op: str) -> float:
    price = costs()
    kind, _, letters = op.partition(":")
    if kind == "sub":
        letter = price.get(op, price.get("sub:" + letters[0], price["ins:"]))
        return price[kind] + letter
    if kind == "ins":
        return price[kind] + price.get(op, price["ins:"])
    return price[kind]


def split_values(cell: str) -> list[str]:
    return [part.strip() for part in cell.split(";") if part.strip()]


def linked_values(cell: str) -> list[str]:
    """«5517160: М0; М1 | 5517205: М0; М1» → значения эталонных характеристик."""
    found = []
    for part in cell.split(" | "):
        found += split_values(part.split(":", 1)[-1])
    return found


def count_spoils(
    value: str, linked: tuple[str, ...], times: int, words_seen: Counter, edits: Counter
) -> None:
    folded = value.casefold()
    best = max(linked, key=lambda v: SequenceMatcher(None, folded, v.casefold()).ratio())
    if SequenceMatcher(None, folded, best.casefold()).ratio() < SAME_VALUE:
        return
    meant = [word["text"] for word in words(best)]
    hit = set()
    for word in words(value):
        kind, ops, index = None, [], -1
        for i, target in enumerate(meant):
            found, found_ops = spoil_kind(word, target)
            if found and (kind is None or RANK[found] < RANK[kind]):
                kind, ops, index = found, found_ops, i
        words_seen[kind or "noise"] += times
        for op in ops:
            edits[op] += times
        hit.add(index)
    kept = len(hit - {-1})
    words_seen["kept"] += kept * times
    words_seen["omit"] += (len(meant) - kept) * times


RANK = {"clean": 0, "layout": 1, "form": 2, "edit": 3}


def spoil_kind(word: Word, target: str) -> tuple[str | None, list[str]]:
    """Правки формы слова считаем как пару окончаний «ый>ая», правки строки по видам."""
    if word["text"] == target:
        return ("layout" if word["fix"] else "clean"), []
    if word["stem"] == stem(target):
        return "form", ["form:" + tails(word["text"], target)]
    distance, ops = edit(word["text"], target, unit)
    if distance <= 2 and len(target) >= 4:
        return "edit", ops
    return None, []


def tails(said: str, meant: str) -> str:
    common = 0
    while common < min(len(said), len(meant)) and said[common] == meant[common]:
        common += 1
    return f"{meant[common:]}>{said[common:]}"


def unit(_: str) -> float:
    return 1.0


def edit(said: str, meant: str, price: Callable[[str], float]) -> tuple[float, list[str]]:
    """Цена правок, которыми meant превращается в said: замена, пропуск, вставка, перестановка
    соседних букв (Дамерау). Возвращает и сами правки, по ним считаются частоты."""
    if abs(len(said) - len(meant)) > MAX_EDITS:
        return inf, []
    rows, cols = len(meant) + 1, len(said) + 1
    table: list[list[tuple[float, tuple[str, ...]]]] = [[(0.0, ())] * cols for _ in range(rows)]
    for i in range(1, rows):
        table[i][0] = step(table[i - 1][0], f"del:{meant[i - 1]}", price)
    for j in range(1, cols):
        table[0][j] = step(table[0][j - 1], f"ins:{said[j - 1]}", price)
    for i in range(1, rows):
        for j in range(1, cols):
            steps = [
                step(table[i - 1][j], f"del:{meant[i - 1]}", price),
                step(table[i][j - 1], f"ins:{said[j - 1]}", price),
            ]
            if meant[i - 1] == said[j - 1]:
                steps.append(table[i - 1][j - 1])
            else:
                steps.append(step(table[i - 1][j - 1], f"sub:{meant[i - 1]}{said[j - 1]}", price))
            if i > 1 and j > 1 and meant[i - 1] == said[j - 2] and meant[i - 2] == said[j - 1]:
                steps.append(step(table[i - 2][j - 2], "swap", price))
            table[i][j] = min(steps, key=lambda found: found[0])
    cost, ops = table[-1][-1]
    return cost, list(ops)


def step(
    before: tuple[float, tuple[str, ...]], op: str, price: Callable[[str], float]
) -> tuple[float, tuple[str, ...]]:
    return before[0] + price(op), (*before[1], op)


@cache
def word_cost(said: str, layout: bool, meant: str) -> tuple[float, FixKind | None]:
    """Самая дешёвая порча, которой слово КТРУ meant стало словом текста said."""
    price = costs()
    ways: list[tuple[float, FixKind | None]] = [
        (price["edit"] + edit(said, meant, op_price)[0], "typo"),
        (price["form"] + price["edit"] + edit(stem(said), stem(meant), op_price)[0], "typo"),
    ]
    if said == meant:
        ways.append((price["clean"], None))
    if stem(said) == stem(meant) and "form:" + tails(said, meant) in price:
        ways.append((price["form"], None))
    cost, kind = min(ways, key=lambda way: way[0])
    if layout:
        return cost + price["layout"] - price["clean"], "layout"
    return cost, kind


def value_prior(positions: list[Position]) -> dict[tuple[str, str], float]:
    """−log доли кандидатов, у которых есть это значение, среди тех, у кого есть характеристика."""
    having: Counter[str] = Counter()
    valued: Counter[tuple[str, str]] = Counter()
    for position in positions:
        for char in position["characteristics"]:
            having[char["name"]] += 1
            for value in {v for v in char["values"] if isinstance(v, str)}:
                valued[char["name"], value] += 1
    return {key: -log(count / having[key[0]]) for key, count in valued.items()}


def explain(
    fragment_words: list[Word],
    options: list[Option],
    prior: dict[tuple[str, str], float],
    product: set[str],
    multiple: set[str],
) -> list[Explained]:
    """Объясняем текст целиком: набор значений с наибольшим суммарным выигрышем над шумом,
    где слово не объясняется дважды и у характеристики одно значение, кроме multiple.
    Ответ даём, если доля веса наборов с этим значением не меньше 1/2: пустой набор
    и наборы-соперники это исходы «текст ни при чём» и «текст про другое»."""
    taken = {i for i, word in enumerate(fragment_words) if word["stem"] in product}
    choice = log(len(options)) if options else 0.0
    alone = []
    for option in options:
        match = score(fragment_words, option, taken, set(), prior, choice)
        if match:
            alone.append(match)
    alone = sorted(alone, key=lambda m: m["gain"], reverse=True)[:MAX_OPTIONS]
    sets = feasible(fragment_words, alone, taken, prior, choice, multiple)
    best_total, best = max(sets, key=lambda found: found[0])
    whole = sum(exp(total - best_total) for total, _ in sets)
    found = []
    for match in best:
        option = match["option"]
        weight = 0.0
        for total, chosen_set in sets:
            if any(other["option"] is option for other in chosen_set):
                weight += exp(total - best_total)
        if weight / whole >= 0.5:
            found.append(match)
    return found


def feasible(
    fragment_words: list[Word],
    alone: list[Explained],
    taken: set[int],
    prior: dict[tuple[str, str], float],
    choice: float,
    multiple: set[str],
) -> list[tuple[float, list[Explained]]]:
    """Все допустимые наборы с выигрышем. Значение в наборе ищет свои слова и слова названия
    среди не объяснённых другими: одно слово может повторяться в тексте для двух значений."""
    sets: list[tuple[float, list[Explained]]] = []

    def grow(start: int, total: float, chosen_set: list[Explained]) -> None:
        sets.append((total, chosen_set))
        names = {m["option"]["char"]["name"] for m in chosen_set}
        explained = set().union(*(m["used"] | m["named"] for m in chosen_set))
        for i in range(start, len(alone)):
            option = alone[i]["option"]
            if option["char"]["name"] in names and option["char"]["name"] not in multiple:
                continue
            match = score(fragment_words, option, taken | explained, explained, prior, choice)
            if match:
                grow(i + 1, total + match["gain"], [*chosen_set, match])

    grow(0, 0.0, [])
    return sets


def score(
    fragment_words: list[Word],
    option: Option,
    taken: set[int],
    values: set[int],
    prior: dict[tuple[str, str], float],
    choice: float,
) -> Explained | None:
    """Выигрыш значения над шумом: −log P(z | шум) − (−log P(z | a) − log P(a)), где P(a)
    это выбор одного из значений кандидатов и доля кандидатов с этим значением."""
    price = costs()
    meant = [word["text"] for word in words(searched_text(option["char"], option["value"]))]
    if not meant:
        return None
    gain = -choice - prior.get((option["char"]["name"], option["value"]), 0.0)
    used: set[int] = set()
    kinds = []
    for target in meant:
        place, cost, kind = closest(fragment_words, target, taken | used)
        if place is None:
            gain -= price["omit"]
            continue
        used.add(place)
        gain += noise() - cost
        kinds.append(kind)
    if not used:
        return None
    named: set[int] = set()
    if option["value"] != "Да":
        for target in [word["text"] for word in words(option["char"]["name"])]:
            place, cost, _ = closest(fragment_words, target, values | used | named)
            if place is not None:
                named.add(place)
                gain += noise() - cost
        if option["source"] == "portal" and not named:
            return None
    if gain <= 0:
        return None
    kind = "layout" if "layout" in kinds else "typo" if "typo" in kinds else None
    return {
        "option": option,
        "gain": gain,
        "used": used,
        "named": named,
        "kind": kind,
    }


def closest(
    fragment_words: list[Word], target: str, busy: set[int]
) -> tuple[int | None, float, FixKind | None]:
    """Слово текста, которое дешевле шума объясняется словом target."""
    best: tuple[int | None, float, FixKind | None] = (None, noise(), None)
    for i, word in enumerate(fragment_words):
        if i in busy:
            continue
        cost, kind = word_cost(word["text"], word["fix"] == "layout", target)
        if cost < best[1]:
            best = (i, cost, kind)
    return best


def options_of(chars: list[Characteristic], source: Source) -> list[Option]:
    seen = set()
    found: list[Option] = []
    for char in chars:
        for value in char["values"]:
            if isinstance(value, str) and value != "Нет" and (char["name"], value) not in seen:
                seen.add((char["name"], value))
                found.append({"char": char, "value": value, "source": source})
    return found


def chosen(
    explained: list[Explained],
    item: Fragment,
    fragment_words: list[Word],
    rule: Rule,
    multiple: set[str],
) -> list[FoundChar]:
    """Одно значение на характеристику, несколько только там, где КТРУ разрешает выбор."""
    found: list[FoundChar] = []
    for match in explained:
        option = match["option"]
        name = option["char"]["name"]
        if any(f["name"] == name for f in found) and name not in multiple:
            continue
        start = min(fragment_words[i]["start"] for i in match["used"])
        end = max(fragment_words[i]["end"] for i in match["used"])
        char = found_char(
            option["char"],
            option["value"],
            option["source"],
            rule,
            item["text"][start:end],
            item["start"] + start,
            item["start"] + end,
        )
        if match["kind"]:
            mark(char, match["kind"])
        found.append(char)
    return found


def found_char(
    char: Characteristic,
    value: str | int | float,
    source: Source,
    rule: Rule,
    fragment: str,
    start: int = 0,
    end: int = 0,
    unit: Unit | None = None,
) -> FoundChar:
    return {
        "name": char["name"],
        "value": value,
        "unit": unit,
        "source": source,
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": char["required"],
        "fix": None,
        "kind": None,
        "rule": rule,
    }


def as_ktru(char: PortalChar) -> Characteristic:
    """Характеристика портала в виде КТРУ, чтобы искать и переводить единицы одним кодом."""
    numeric = "числов" in char["type"].lower()
    item: Characteristic = {
        "code": "",
        "name": char["name"],
        "type": "number" if numeric else "text",
        "selection": "single",
        "required": False,
        "values": [] if numeric else list(char["values"]),
    }
    unit = (char["unit"] or "").strip("() ")
    if unit:
        item["unit"] = {"code": UNIT_ALIASES.get(unit.lower(), unit), "name": unit}
    return item


def allowed_values(chars: list[Characteristic]) -> dict[str, list[Value]]:
    """Значения характеристики по всем кандидатам: у соседних кодов списки разные."""
    allowed: dict[str, list[Value]] = {}
    for char in chars:
        allowed.setdefault(char["name"], []).extend(char["values"])
    return allowed


def multiple_names(chars: list[Characteristic]) -> set[str]:
    return {char["name"] for char in chars if char["selection"] == "multiple"}


def told_pairs(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Человек дописывает исправление в конец, поэтому из явных значений верны последние."""
    last: dict[str, list[FoundChar]] = {}
    for item in fragments:
        if item["form"] == "pair":
            found = pair_chars(item, chars, portal_chars, allowed, product, prior)
            if found:
                last[found[0]["name"]] = found
    return [char for found in last.values() for char in found]


def pair_chars(
    item: Fragment,
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    offset = name_offset(item["name"] or "", product)
    if offset is None:
        return []
    name = (item["name"] or "")[offset:]
    start = item["start"] + offset
    fragment = item["text"][offset:]
    char = named(chars, name)
    source: Source = "ktru"
    if char is None:
        char = named(portal_chars, name)
        source = "portal"
    if char is None:
        return [new_char(name, item["value"], fragment, start, item["end"])]
    values = allowed.get(char["name"], []) if source == "ktru" else char["values"]
    if char["type"] == "number":
        if item["number"] is None:
            return []
        found = [number_char(char, source, "pair", item["number"], item["unit"], values, fragment)]
    else:
        whole = char.copy()
        whole["values"] = values
        found = text_pair(whole, source, item["value"], fragment, prior, multiple_names(chars))
    for char_found in found:
        char_found["start"], char_found["end"] = start, item["end"]
    return found


def name_offset(name: str, product: set[str]) -> int | None:
    """«Джемпер рост: 170»: слова товара перед названием характеристики к нему не относятся."""
    for word in words(name):
        if word["stem"] not in product:
            return word["start"]
    return None


def named(chars: list[Characteristic], name: str) -> Characteristic | None:
    """Характеристика с тем же названием, иначе та, в названии которой есть все слова."""
    for char in chars:
        if char["name"].casefold() == name.casefold():
            return char
    wanted = stems(name)
    for char in chars:
        if wanted and wanted <= stems(char["name"]):
            return char
    return None


def new_char(name: str, raw: str, fragment: str, start: int, end: int) -> FoundChar:
    return {
        "name": name[:1].upper() + name[1:],
        "value": raw,
        "unit": None,
        "source": "new",
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": False,
        "fix": None,
        "kind": None,
        "rule": "pair",
    }


def text_pair(
    char: Characteristic,
    source: Source,
    raw: str,
    fragment: str,
    prior: dict[tuple[str, str], float],
    multiple: set[str],
) -> list[FoundChar]:
    """Значение пары это те значения списка, к которым текст приводится дешевле шума. Списки
    портала собраны из того, что вводили поставщики, поэтому «нет в списке» только про КТРУ."""
    options = {v.casefold(): v for v in char["values"] if isinstance(v, str)}
    if raw.casefold() in options:
        return [found_char(char, options[raw.casefold()], source, "pair", fragment)]
    item: Fragment = {
        "form": "words",
        "name": None,
        "value": raw,
        "text": raw,
        "number": None,
        "unit": None,
        "start": 0,
        "end": len(raw),
    }
    raw_words = words(raw)
    pair_options = [o for o in options_of([char], source) if o["value"] != "Да"]
    found = chosen(
        explain(raw_words, pair_options, prior, set(), multiple), item, raw_words, "pair", multiple
    )
    if found:
        for char_found in found:
            char_found["fragment"] = fragment
        return found
    left = found_char(char, raw, source, "pair", fragment)
    if source == "ktru":
        shown = ", ".join(sorted(set(options.values()))[:SHOWN_OPTIONS])
        mark(left, "not_in_list", options=shown)
    return [left]


def told_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    product: set[str],
    shared: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Значения без названия: «деревянный» → Тип каркаса: Деревянный. Значение портала берём,
    только если рядом названа и сама характеристика: в списках портала много мусора."""
    ktru = options_of(chars, "ktru")
    portal = options_of(portal_chars, "portal")
    multiple = multiple_names(chars)
    found = []
    for item in fragments:
        if item["form"] == "words":
            fragment_words = words(item["text"])
            found += chosen(
                explain(fragment_words, ktru, prior, shared, multiple),
                item,
                fragment_words,
                "value",
                multiple,
            )
            # Слова товара не подсказывают характеристику портала: «Вид стола: Письменный».
            told = [word for word in fragment_words if word["stem"] not in product]
            found += chosen(
                explain(told, portal, prior, shared, multiple), item, told, "value", multiple
            )
    return found


def searched_text(char: Characteristic, value: str) -> str:
    """«Да» ищем по названию: «из натуральной кожи» значит «Натуральная кожа: Да»."""
    if value == "Да":
        return char["name"]
    return value


def number_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
) -> list[FoundChar]:
    """Числа в единице КТРУ: «2,2 см» → Толщина столешницы 22 мм, «высота 75» → 75 см."""
    numeric = numeric_chars(chars)
    portal_numeric = numeric_chars(portal_chars)
    found: list[FoundChar] = []
    for item in fragments:
        if item["form"] != "number" or item["number"] is None:
            continue
        name_stems = stems(item["name"] or "")
        char = owner(numeric, item["unit"], name_stems, guess=True)
        source: Source = "ktru"
        if char is None:
            # В названиях портала много мусора, поэтому название должно быть в тексте целиком.
            told = name_stems - product
            named_here = [c for c in portal_numeric if stems(c["name"]) <= told]
            char = owner(named_here, item["unit"], told, guess=False)
            source = "portal"
        if char is None:
            continue
        values = allowed.get(char["name"], [])
        char_found = number_char(
            char, source, "number", item["number"], item["unit"], values, item["text"]
        )
        char_found["start"], char_found["end"] = item["start"], item["end"]
        found.append(char_found)
    return found


def numeric_chars(chars: list[Characteristic]) -> list[Characteristic]:
    numeric = {}
    for char in chars:
        if char["type"] == "number" and "unit" in char:
            numeric[char["name"]] = char
    return list(numeric.values())


def number_char(
    char: Characteristic,
    source: Source,
    rule: Rule,
    number: float,
    code: str | None,
    allowed: list[Value],
    fragment: str,
) -> FoundChar:
    unit: Unit | None = char["unit"] if "unit" in char else None
    value = number
    if unit and code:
        value = convert(number, code, unit["code"])
    item = found_char(char, tidy(value), source, rule, fragment, unit=unit)
    if unit and code and code != unit["code"]:
        mark(item, "unit")
    check_range(item, allowed)
    return item


def check_range(item: FoundChar, allowed: list[Value]) -> None:
    """Число вне коридора КТРУ: правим, если сдвиг на порядок его объясняет, иначе помечаем."""
    corridor = bounds(allowed)
    value = item["value"]
    if corridor is None or isinstance(value, str) or within(value, corridor):
        return
    for shift in SHIFTS:
        if within(value * shift, corridor):
            item["value"] = tidy(value * shift)
            mark(item, "wrong_unit")
            return
    low, high = corridor
    mark(item, "out_of_range", low=show(low), high=show(high))


def within(value: float, corridor: tuple[float, float]) -> bool:
    return corridor[0] <= value <= corridor[1]


def owner(
    chars: list[Characteristic], code: str | None, name_stems: set[str], guess: bool
) -> Characteristic | None:
    """Характеристика той же величины, чьё название ближе к словам перед числом. Угадывать
    без подсказки в тексте можно, только если выбора нет, единица названа и guess разрешён."""
    best = None
    best_hits = 0
    same_dimension = 0
    for char in chars:
        if code is not None and scale(char["unit"]["code"])[0] != scale(code)[0]:
            continue
        same_dimension += 1
        hits = len(stems(char["name"]) & name_stems)
        if best is None or hits > best_hits:
            best = char
            best_hits = hits
    if best_hits == 0 and (not guess or same_dimension > 1 or code is None):
        return None
    return best
