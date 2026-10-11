"""Правку помечаем, только если калиброванная P(порча | текст) по всем порчам не ниже F*/(1 + β²).

Вывод: судья на dev 0.414 → 0.253 (верно 554, отказ 85, неверно 44; у noisy_channel было
0.262). В одном прогоне с noisy_channel верно 539 → 541, неверно 46 → 44, лучше 4, хуже 2.
Правки: ложные на чистых 9 → 6, пропущенные 14 → 15, F0.5 0.907 → 0.926. Ушли ложные
«опечатки» у форм слова («из металла», «двухсторонное», «металлическые»), появилась верная
«ножи → Ножки». Пропущены опечатки, похожие на форму («деревяная», «треугольнаая»).
Порог 0.79 при F* = 0.985 на синтетике; на ней же F0.5 при этом пороге 0.982, лучший 0.985:
формула Lipton и др. работает, когда вероятность откалибрована. Одной температуры
не хватило, нужен сдвиг Платта.

Значение выбирается как в noisy_channel: самая дешёвая порча и отказ по Чоу. Пометка своя:
P(порча | слово) это сумма по путям «правка строки» и «правка после формы» против «чисто»
и «форма слова», где форма это любая пара окончаний, а пара, которой не было в листах 2,
стоит по Гуду и Тьюрингу. Для значения P = 1 − Π(1 − P слова). Вероятность калибруется
по Платту на синтетике: значения КТРУ, слова которых чисты, меняют форму или получают
правку с частотами листов 2. Порог следует из F0.5 (Lipton, Elkan, Narayanaswamy 2014):
добавить пометку с калиброванной вероятностью p выгодно, только если p ≥ F*/(1 + β²).
"""

import json
from collections import Counter
from collections.abc import Callable
from difflib import SequenceMatcher
from functools import cache
from math import exp, inf, log
from pathlib import Path
from random import Random
from typing import Literal, TypedDict

from openpyxl import load_workbook

from contracts.pipeline import Candidate, FixKind, FoundChar, Fragment, Rule
from contracts.ref import Characteristic, Portal, PortalChar, Position, Unit, Value
from pipeline.explain import mark
from pipeline.run import reference
from pipeline.text import Word, stem, stems, words
from pipeline.values import UNIT_ALIASES, bounds, convert, same, scale, show, tidy
from scripts.products import RAW_DIR, SHEETS, products

# Перепутанная единица чаще всего сдвигает число на порядок: см вместо мм, м вместо см.
SHIFTS = (10, 0.1, 100, 0.01, 1000, 0.001)
SHOWN_OPTIONS = 6
COSTS_PATH = Path("data/processed/link_noisy_channel_costs.json")
# Написание базовой характеристики сравниваем с эталонным, только если это одно значение.
SAME_VALUE = 0.5
MAX_EDITS = 3
RUSSIAN = "абвгдежзийклмнопрстуфхцчшщъыьэюя"
ALPHABET = len(RUSSIAN) + 1
OPS = ("sub", "del", "ins", "swap")
# В правках точность важнее полноты: ложная правка чистого значения хуже пропущенной.
BETA = 0.5
CALIBRATION_VALUES = 20000

type Source = Literal["ktru", "portal"]


class Option(TypedDict):
    char: Characteristic
    value: str
    source: Source


class Explained(TypedDict):
    option: Option
    gain: float
    used: set[int]
    named: set[int]
    kind: FixKind | None


def link(
    fragments: list[Fragment],
    candidates: list[Candidate],
    section: str | None,
    ref: list[Position],
    portal: Portal,
) -> list[FoundChar]:
    by_code = {p["code"]: p for p in ref}
    positions = [by_code[candidate["code"]] for candidate in candidates]
    chars = [char for position in positions for char in position["characteristics"]]
    allowed = allowed_values(chars)
    portal_chars = [as_ktru(char) for char in portal.get(section or "", [])]
    product = stems(" ".join(c["name"] for c in candidates))
    prior = value_prior(positions)

    pairs = told_pairs(fragments, chars, portal_chars, allowed, product, prior)
    shared = set.intersection(*(stems(c["name"]) for c in candidates)) if candidates else set()
    values = told_values(fragments, chars, portal_chars, product, shared, prior)
    numbers = number_values(fragments, chars, portal_chars, allowed, product)
    found = choose(pairs, values, numbers)
    # Обязательность у соседних кодов разная: обязательна, если обязательна хоть у одного.
    required = {char["name"] for char in chars if char["required"]}
    for char in found:
        char["required"] = char["source"] == "ktru" and char["name"] in required
    return found


def choose(
    pairs: list[FoundChar], values: list[FoundChar], numbers: list[FoundChar]
) -> list[FoundChar]:
    """Явная пара сильнее значения без названия, оно сильнее числа. Значения одного правила
    уже выбраны каналом, между правилами разные значения дают правку «два значения»."""
    by_name: dict[str, list[FoundChar]] = {}
    for rule in (pairs, values, numbers):
        mine: dict[str, list[FoundChar]] = {}
        for char in rule:
            mine.setdefault(char["name"], []).append(char)
        for name, chars in mine.items():
            kept = by_name.get(name)
            if kept is None:
                by_name[name] = chars
                continue
            other = [c for c in chars if not any(same(k["value"], c["value"]) for k in kept)]
            if other:
                mark(kept[0], "conflict", other=other[0]["fragment"])
    return [char for chars in by_name.values() for char in chars]


@cache
def costs() -> dict[str, float]:
    """−log частоты каждой порчи. Счётчики собираются один раз по листам 2 в data/processed."""
    if not COSTS_PATH.exists():
        COSTS_PATH.write_text(json.dumps(mined_counts(), ensure_ascii=False, indent=1))
    counts = json.loads(COSTS_PATH.read_text())
    return priced(Counter(counts["words"]), Counter(counts["edits"]))


@cache
def noise() -> float:
    """Слово шума это порча «лишнее слово» и любое слово из тех, что поставщики пишут в товарах
    листа 1, на выбор."""
    vocabulary = set()
    for product in products({}, build=False):
        vocabulary |= {word["text"] for word in words(product["name"])}
        for name, value in product["pairs"]:
            vocabulary |= {word["text"] for word in words(f"{name} {value}")}
    return costs()["noise"] + log(len(vocabulary))


def mined_counts() -> dict[str, Counter[str]]:
    written: Counter[tuple[str, tuple[str, ...]]] = Counter()
    for sheet in SHEETS:
        book = load_workbook(RAW_DIR / sheet, read_only=True)
        for row in book.worksheets[1].iter_rows(min_row=2, max_col=11, values_only=True):
            linked = tuple(linked_values(str(row[10] or "")))
            if row[1] and row[6] and linked:
                for value in split_values(str(row[6])):
                    written[value, linked] += 1
        book.close()
    words_seen: Counter[str] = Counter()
    edits: Counter[str] = Counter()
    for (value, linked), times in written.items():
        count_spoils(value, linked, times, words_seen, edits)
    return {"words": words_seen, "edits": edits}


def priced(words_seen: Counter[str], edits: Counter[str]) -> dict[str, float]:
    """Цена вида порчи по словам, цена правки по виду и по букве: P(буква текста | буква КТРУ)
    для замены и P(буква) для вставки, со сглаживанием на букву алфавита."""
    said = sum(words_seen[k] for k in ("clean", "layout", "form", "edit", "noise"))
    found = {}
    for kind in ("clean", "layout", "form", "edit", "noise"):
        found[kind] = -log((words_seen[kind] + 1) / said)
    found["omit"] = -log((words_seen["omit"] + 1) / (words_seen["kept"] + words_seen["omit"]))
    by_kind: Counter[str] = Counter()
    by_letter: Counter[str] = Counter()
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        by_kind[kind] += count
        if kind == "sub":
            by_letter["sub:" + letters[0]] += count
    total = sum(by_kind[kind] for kind in OPS) + len(OPS)
    for kind in OPS:
        found[kind] = -log((by_kind[kind] + 1) / total)
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        if kind == "sub":
            found[op] = -log((count + 1) / (by_letter["sub:" + letters[0]] + ALPHABET))
        elif kind == "ins":
            found[op] = -log((count + 1) / (by_kind["ins"] + ALPHABET))
        elif kind == "form":
            found[op] = 0.0
    for letter in RUSSIAN:
        found.setdefault("sub:" + letter, -log(1 / (by_letter["sub:" + letter] + ALPHABET)))
    found["ins:"] = -log(1 / (by_kind["ins"] + ALPHABET))
    return found


def op_price(op: str) -> float:
    price = costs()
    kind, _, letters = op.partition(":")
    if kind == "sub":
        letter = price.get(op, price.get("sub:" + letters[0], price["ins:"]))
        return price[kind] + letter
    if kind == "ins":
        return price[kind] + price.get(op, price["ins:"])
    return price[kind]


def split_values(cell: str) -> list[str]:
    return [part.strip() for part in cell.split(";") if part.strip()]


def linked_values(cell: str) -> list[str]:
    """«5517160: М0; М1 | 5517205: М0; М1» → значения эталонных характеристик."""
    found = []
    for part in cell.split(" | "):
        found += split_values(part.split(":", 1)[-1])
    return found


def count_spoils(
    value: str, linked: tuple[str, ...], times: int, words_seen: Counter, edits: Counter
) -> None:
    folded = value.casefold()
    best = max(linked, key=lambda v: SequenceMatcher(None, folded, v.casefold()).ratio())
    if SequenceMatcher(None, folded, best.casefold()).ratio() < SAME_VALUE:
        return
    meant = [word["text"] for word in words(best)]
    hit = set()
    for word in words(value):
        kind, ops, index = None, [], -1
        for i, target in enumerate(meant):
            found, found_ops = spoil_kind(word, target)
            if found and (kind is None or RANK[found] < RANK[kind]):
                kind, ops, index = found, found_ops, i
        words_seen[kind or "noise"] += times
        for op in ops:
            edits[op] += times
        hit.add(index)
    kept = len(hit - {-1})
    words_seen["kept"] += kept * times
    words_seen["omit"] += (len(meant) - kept) * times


RANK = {"clean": 0, "layout": 1, "form": 2, "edit": 3}


def spoil_kind(word: Word, target: str) -> tuple[str | None, list[str]]:
    """Правки формы слова считаем как пару окончаний «ый>ая», правки строки по видам."""
    if word["text"] == target:
        return ("layout" if word["fix"] else "clean"), []
    if word["stem"] == stem(target):
        return "form", ["form:" + tails(word["text"], target)]
    distance, ops = edit(word["text"], target, unit)
    if distance <= 2 and len(target) >= 4:
        return "edit", ops
    return None, []


def tails(said: str, meant: str) -> str:
    common = 0
    while common < min(len(said), len(meant)) and said[common] == meant[common]:
        common += 1
    return f"{meant[common:]}>{said[common:]}"


def unit(_: str) -> float:
    return 1.0


def edit(said: str, meant: str, price: Callable[[str], float]) -> tuple[float, list[str]]:
    """Цена правок, которыми meant превращается в said: замена, пропуск, вставка, перестановка
    соседних букв (Дамерау). Возвращает и сами правки, по ним считаются частоты."""
    if abs(len(said) - len(meant)) > MAX_EDITS:
        return inf, []
    rows, cols = len(meant) + 1, len(said) + 1
    table: list[list[tuple[float, tuple[str, ...]]]] = [[(0.0, ())] * cols for _ in range(rows)]
    for i in range(1, rows):
        table[i][0] = step(table[i - 1][0], f"del:{meant[i - 1]}", price)
    for j in range(1, cols):
        table[0][j] = step(table[0][j - 1], f"ins:{said[j - 1]}", price)
    for i in range(1, rows):
        for j in range(1, cols):
            steps = [
                step(table[i - 1][j], f"del:{meant[i - 1]}", price),
                step(table[i][j - 1], f"ins:{said[j - 1]}", price),
            ]
            if meant[i - 1] == said[j - 1]:
                steps.append(table[i - 1][j - 1])
            else:
                steps.append(step(table[i - 1][j - 1], f"sub:{meant[i - 1]}{said[j - 1]}", price))
            if i > 1 and j > 1 and meant[i - 1] == said[j - 2] and meant[i - 2] == said[j - 1]:
                steps.append(step(table[i - 2][j - 2], "swap", price))
            table[i][j] = min(steps, key=lambda found: found[0])
    cost, ops = table[-1][-1]
    return cost, list(ops)


def step(
    before: tuple[float, tuple[str, ...]], op: str, price: Callable[[str], float]
) -> tuple[float, tuple[str, ...]]:
    return before[0] + price(op), (*before[1], op)


@cache
def word_cost(said: str, layout: bool, meant: str) -> tuple[float, FixKind | None]:
    """Самая дешёвая порча, которой слово КТРУ meant стало словом текста said."""
    price = costs()
    ways: list[tuple[float, FixKind | None]] = [
        (price["edit"] + edit(said, meant, op_price)[0], "typo"),
        (price["form"] + price["edit"] + edit(stem(said), stem(meant), op_price)[0], "typo"),
    ]
    if said == meant:
        ways.append((price["clean"], None))
    if stem(said) == stem(meant) and "form:" + tails(said, meant) in price:
        ways.append((price["form"], None))
    cost, kind = min(ways, key=lambda way: way[0])
    if layout:
        return cost + price["layout"] - price["clean"], "layout"
    return cost, kind


@cache
def fix_chance(said: str, meant: str) -> float:
    """P(слово испорчено правкой строки | said, meant): сумма по всем порчам, а не лучшая.
    Форма слова это любая замена окончания, а не только при совпавшей грубой основе:
    «пластика» и «пластик» основы дают разные, и форма выходила опечаткой."""
    price = costs()
    clean = []
    spoiled = []
    if said == meant:
        clean.append(price["clean"])
    elif said[:1] == meant[:1]:
        clean.append(price["form"] + tail_cost(tails(said, meant)))
    typo, ops = edit(said, meant, op_price)
    if ops:
        spoiled.append(price["edit"] + typo)
    typo, ops = edit(stem(said), stem(meant), op_price)
    if ops:
        spoiled.append(price["form"] + price["edit"] + typo)
    bad = sum(exp(-cost) for cost in spoiled)
    good = sum(exp(-cost) for cost in clean)
    return bad / (bad + good) if bad + good else 1.0


def tail_cost(tail: str) -> float:
    """−log P(пара окончаний | форма). Новая пара по Гуду и Тьюрингу стоит долю пар,
    встреченных один раз, а её окончания выбираются независимо из виденных: иначе любая
    опечатка в конце слова объяснялась бы формой."""
    seen = form_tails()
    total = sum(seen.values())
    if tail in seen:
        return -log(seen[tail] / total)
    once = sum(1 for count in seen.values() if count == 1)
    was, _, now = tail.partition(">")
    return -log(once / total) + ending_cost(was, 0) + ending_cost(now, 2)


def ending_cost(ending: str, side: int) -> float:
    """Окончание, которого не было и среди отдельных окончаний, набирается буква за буквой:
    тогда «стул» вместо «стол» дороже опечатки."""
    seen = form_tails()
    endings: Counter[str] = Counter()
    for pair, count in seen.items():
        endings[pair.partition(">")[side]] += count
    if ending in endings:
        return -log(endings[ending] / sum(seen.values()))
    return len(ending) * log(ALPHABET)


@cache
def form_tails() -> Counter[str]:
    costs()
    counts = json.loads(COSTS_PATH.read_text())["edits"]
    return Counter({op[5:]: n for op, n in counts.items() if op.startswith("form:")})


def value_chance(chances: list[float]) -> float:
    clean = 1.0
    for chance in chances:
        clean *= 1 - chance
    return 1 - clean


def flagged(chances: list[float]) -> bool:
    slope, bias, threshold = calibration()
    return calibrated(value_chance(chances), slope, bias) >= threshold


def log_odds(chance: float) -> float:
    chance = min(max(chance, 1e-12), 1 - 1e-12)
    return log(chance / (1 - chance))


def calibrated(chance: float, slope: float, bias: float) -> float:
    return sigmoid(slope * log_odds(chance) + bias)


def sigmoid(x: float) -> float:
    if x < 0:
        return exp(x) / (1 + exp(x))
    return 1 / (1 + exp(-x))


@cache
def calibration() -> tuple[float, float, float]:
    """Калибровка Платта и порог пометки на синтетике из значений КТРУ, испорченных теми же
    порчами с теми же частотами, что в листах 2. Одной температуры (Guo и др. 2017) мало:
    канал уверен в опечатке и там, где это новая форма слова, нужен ещё и сдвиг. Порог
    F*/(1 + β²) по Lipton и др. 2014: при калиброванной вероятности помечать выгодно,
    только если она не ниже этой доли лучшего достижимого F."""
    chances, labels = synthetic()
    slope, bias = platt([log_odds(c) for c in chances], labels)
    scored = sorted(
        zip((calibrated(c, slope, bias) for c in chances), labels, strict=True), reverse=True
    )
    return slope, bias, best_f(scored) / (1 + BETA**2)


def platt(odds: list[float], labels: list[bool]) -> tuple[float, float]:
    """Наклон и сдвиг логистической кривой по сетке с наименьшей логарифмической потерей.
    Метки смягчены, как у Платта (1999): синтетика почти разделима, и без этого наклон
    уходит в бесконечность."""
    positives = sum(labels)
    high = (positives + 1) / (positives + 2)
    low = 1 / (len(labels) - positives + 2)
    points: Counter[tuple[float, float]] = Counter()
    for z, label in zip(odds, labels, strict=True):
        points[round(z, 2), high if label else low] += 1
    grid = [(exp(a / 10), b / 10) for a in range(-30, 31) for b in range(-150, 151)]
    return min(grid, key=lambda line: platt_loss(points, *line))


def platt_loss(points: Counter[tuple[float, float]], slope: float, bias: float) -> float:
    loss = 0.0
    for (z, target), count in points.items():
        p = min(max(sigmoid(slope * z + bias), 1e-12), 1 - 1e-12)
        loss -= count * (target * log(p) + (1 - target) * log(1 - p))
    return loss


def best_f(scored: list[tuple[float, bool]]) -> float:
    positives = sum(label for _, label in scored)
    best = 0.0
    hits = 0
    for flags, (_, label) in enumerate(scored, 1):
        hits += label
        precision = hits / flags
        recall = hits / positives if positives else 0.0
        if hits:
            f = (1 + BETA**2) * precision * recall / (BETA**2 * precision + recall)
            best = max(best, f)
    return best


def synthetic() -> tuple[list[float], list[bool]]:
    """Каждое слово значения остаётся чистым, меняет форму или получает правку строки
    с частотами листов 2; правка выбирается из правок, найденных там же."""
    rng = Random(0)
    costs()
    counts = json.loads(COSTS_PATH.read_text())
    seen = counts["words"]
    kinds = ("clean", "form", "edit")
    weights = [seen[kind] for kind in kinds]
    values = ktru_values()
    chances = []
    labels = []
    for value in rng.choices(values, k=CALIBRATION_VALUES):
        word_chances = []
        spoiled = False
        for word in words(value):
            meant = word["text"]
            said = meant
            kind = rng.choices(kinds, weights)[0]
            if kind == "form":
                said = inflected(meant, rng)
            if kind == "edit" and len(meant) >= 4:
                said = typed(meant, counts["edits"], rng)
                spoiled = spoiled or said != meant
            word_chances.append(fix_chance(said, meant))
        if word_chances:
            chances.append(value_chance(word_chances))
            labels.append(spoiled)
    return chances, labels


def ktru_values() -> list[str]:
    found = set()
    for position in reference():
        for char in position["characteristics"]:
            for value in char["values"]:
                if isinstance(value, str) and value not in ("Да", "Нет"):
                    found.add(value)
    return sorted(found)


def inflected(meant: str, rng: Random) -> str:
    fitting = []
    weights = []
    for tail, count in form_tails().items():
        was, _, now = tail.partition(">")
        if meant.endswith(was) and len(was) < len(meant):
            fitting.append(meant[: len(meant) - len(was)] + now)
            weights.append(count)
    return rng.choices(fitting, weights)[0] if fitting else meant


def typed(meant: str, edits: dict[str, int], rng: Random) -> str:
    fitting = []
    weights = []
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        if kind in ("sub", "del") and letters[:1] in meant or kind in ("ins", "swap"):
            fitting.append(op)
            weights.append(count)
    kind, _, letters = rng.choices(fitting, weights)[0].partition(":")
    if kind == "ins":
        at = rng.randrange(len(meant) + 1)
        return meant[:at] + letters + meant[at:]
    if kind == "swap":
        at = rng.randrange(len(meant) - 1)
        return meant[:at] + meant[at + 1] + meant[at] + meant[at + 2 :]
    places = [i for i, letter in enumerate(meant) if letter == letters[0]]
    at = rng.choice(places)
    return meant[:at] + letters[1:] + meant[at + 1 :]


def value_prior(positions: list[Position]) -> dict[tuple[str, str], float]:
    """−log доли кандидатов, у которых есть это значение, среди тех, у кого есть характеристика."""
    having: Counter[str] = Counter()
    valued: Counter[tuple[str, str]] = Counter()
    for position in positions:
        for char in position["characteristics"]:
            having[char["name"]] += 1
            for value in {v for v in char["values"] if isinstance(v, str)}:
                valued[char["name"], value] += 1
    return {key: -log(count / having[key[0]]) for key, count in valued.items()}


def explain(
    fragment_words: list[Word],
    options: list[Option],
    prior: dict[tuple[str, str], float],
    product: set[str],
) -> list[Explained]:
    """Жадно объясняем текст: значение с наибольшим выигрышем над шумом забирает свои слова
    и слова названия своей характеристики, остальные пересчитываются на оставшихся. Слова,
    общие для названий всех кандидатов, уже объяснены товаром. Если соперник за те же слова
    почти так же вероятен (P лучшего < 1/2), слова забираем, а ответа не даём."""
    taken = {i for i, word in enumerate(fragment_words) if word["stem"] in product}
    values: set[int] = set()
    found = []
    while True:
        scored = []
        for option in options:
            match = score(fragment_words, option, taken, values, prior, log(len(options)))
            if match:
                scored.append(match)
        if not scored:
            return found
        best = max(scored, key=lambda s: s["gain"])
        rivals = [s for s in scored if s is not best and s["used"] & best["used"]]
        odds = exp(-best["gain"]) + sum(exp(s["gain"] - best["gain"]) for s in rivals)
        taken |= best["used"] | best["named"]
        values |= best["used"]
        if 1 / (1 + odds) >= 0.5:
            found.append(best)


def score(
    fragment_words: list[Word],
    option: Option,
    taken: set[int],
    values: set[int],
    prior: dict[tuple[str, str], float],
    choice: float,
) -> Explained | None:
    """Выигрыш значения над шумом: −log P(z | шум) − (−log P(z | a) − log P(a)), где P(a)
    это выбор одного из значений кандидатов и доля кандидатов с этим значением."""
    price = costs()
    meant = [word["text"] for word in words(searched_text(option["char"], option["value"]))]
    if not meant:
        return None
    gain = -choice - prior.get((option["char"]["name"], option["value"]), 0.0)
    used: set[int] = set()
    kinds = []
    chances = []
    for target in meant:
        place, cost, kind = closest(fragment_words, target, taken | used)
        if place is None:
            gain -= price["omit"]
            continue
        used.add(place)
        gain += noise() - cost
        kinds.append(kind)
        chances.append(fix_chance(fragment_words[place]["text"], target))
    if not used:
        return None
    named: set[int] = set()
    if option["value"] != "Да":
        for target in {word["text"] for word in words(option["char"]["name"])}:
            place, cost, _ = closest(fragment_words, target, values | used | named)
            if place is not None:
                named.add(place)
                gain += noise() - cost
        if option["source"] == "portal" and not named:
            return None
    if gain <= 0:
        return None
    kind = "layout" if "layout" in kinds else "typo" if flagged(chances) else None
    return {"option": option, "gain": gain, "used": used, "named": named, "kind": kind}


def closest(
    fragment_words: list[Word], target: str, busy: set[int]
) -> tuple[int | None, float, FixKind | None]:
    """Слово текста, которое дешевле шума объясняется словом target."""
    best: tuple[int | None, float, FixKind | None] = (None, noise(), None)
    for i, word in enumerate(fragment_words):
        if i in busy:
            continue
        cost, kind = word_cost(word["text"], word["fix"] == "layout", target)
        if cost < best[1]:
            best = (i, cost, kind)
    return best


def options_of(chars: list[Characteristic], source: Source) -> list[Option]:
    seen = set()
    found: list[Option] = []
    for char in chars:
        for value in char["values"]:
            if isinstance(value, str) and value != "Нет" and (char["name"], value) not in seen:
                seen.add((char["name"], value))
                found.append({"char": char, "value": value, "source": source})
    return found


def chosen(
    explained: list[Explained],
    item: Fragment,
    fragment_words: list[Word],
    rule: Rule,
    multiple: set[str],
) -> list[FoundChar]:
    """Одно значение на характеристику, несколько только там, где КТРУ разрешает выбор."""
    found: list[FoundChar] = []
    for match in explained:
        option = match["option"]
        name = option["char"]["name"]
        if any(f["name"] == name for f in found) and name not in multiple:
            continue
        start = min(fragment_words[i]["start"] for i in match["used"])
        end = max(fragment_words[i]["end"] for i in match["used"])
        char = found_char(
            option["char"],
            option["value"],
            option["source"],
            rule,
            item["text"][start:end],
            item["start"] + start,
            item["start"] + end,
        )
        if match["kind"]:
            mark(char, match["kind"])
        found.append(char)
    return found


def found_char(
    char: Characteristic,
    value: str | int | float,
    source: Source,
    rule: Rule,
    fragment: str,
    start: int = 0,
    end: int = 0,
    unit: Unit | None = None,
) -> FoundChar:
    return {
        "name": char["name"],
        "value": value,
        "unit": unit,
        "source": source,
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": char["required"],
        "fix": None,
        "kind": None,
        "rule": rule,
    }


def as_ktru(char: PortalChar) -> Characteristic:
    """Характеристика портала в виде КТРУ, чтобы искать и переводить единицы одним кодом."""
    numeric = "числов" in char["type"].lower()
    item: Characteristic = {
        "code": "",
        "name": char["name"],
        "type": "number" if numeric else "text",
        "selection": "single",
        "required": False,
        "values": [] if numeric else list(char["values"]),
    }
    unit = (char["unit"] or "").strip("() ")
    if unit:
        item["unit"] = {"code": UNIT_ALIASES.get(unit.lower(), unit), "name": unit}
    return item


def allowed_values(chars: list[Characteristic]) -> dict[str, list[Value]]:
    """Значения характеристики по всем кандидатам: у соседних кодов списки разные."""
    allowed: dict[str, list[Value]] = {}
    for char in chars:
        allowed.setdefault(char["name"], []).extend(char["values"])
    return allowed


def multiple_names(chars: list[Characteristic]) -> set[str]:
    return {char["name"] for char in chars if char["selection"] == "multiple"}


def told_pairs(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Человек дописывает исправление в конец, поэтому из явных значений верны последние."""
    last: dict[str, list[FoundChar]] = {}
    for item in fragments:
        if item["form"] == "pair":
            found = pair_chars(item, chars, portal_chars, allowed, product, prior)
            if found:
                last[found[0]["name"]] = found
    return [char for found in last.values() for char in found]


def pair_chars(
    item: Fragment,
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    offset = name_offset(item["name"] or "", product)
    if offset is None:
        return []
    name = (item["name"] or "")[offset:]
    start = item["start"] + offset
    fragment = item["text"][offset:]
    char = named(chars, name)
    source: Source = "ktru"
    if char is None:
        char = named(portal_chars, name)
        source = "portal"
    if char is None:
        return [new_char(name, item["value"], fragment, start, item["end"])]
    values = allowed.get(char["name"], []) if source == "ktru" else char["values"]
    if char["type"] == "number":
        if item["number"] is None:
            return []
        found = [number_char(char, source, "pair", item["number"], item["unit"], values, fragment)]
    else:
        whole = char.copy()
        whole["values"] = values
        found = text_pair(whole, source, item["value"], fragment, prior, multiple_names(chars))
    for char_found in found:
        char_found["start"], char_found["end"] = start, item["end"]
    return found


def name_offset(name: str, product: set[str]) -> int | None:
    """«Джемпер рост: 170»: слова товара перед названием характеристики к нему не относятся."""
    for word in words(name):
        if word["stem"] not in product:
            return word["start"]
    return None


def named(chars: list[Characteristic], name: str) -> Characteristic | None:
    """Характеристика с тем же названием, иначе та, в названии которой есть все слова."""
    for char in chars:
        if char["name"].casefold() == name.casefold():
            return char
    wanted = stems(name)
    for char in chars:
        if wanted and wanted <= stems(char["name"]):
            return char
    return None


def new_char(name: str, raw: str, fragment: str, start: int, end: int) -> FoundChar:
    return {
        "name": name[:1].upper() + name[1:],
        "value": raw,
        "unit": None,
        "source": "new",
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": False,
        "fix": None,
        "kind": None,
        "rule": "pair",
    }


def text_pair(
    char: Characteristic,
    source: Source,
    raw: str,
    fragment: str,
    prior: dict[tuple[str, str], float],
    multiple: set[str],
) -> list[FoundChar]:
    """Значение пары это те значения списка, к которым текст приводится дешевле шума. Списки
    портала собраны из того, что вводили поставщики, поэтому «нет в списке» только про КТРУ."""
    options = {v.casefold(): v for v in char["values"] if isinstance(v, str)}
    if raw.casefold() in options:
        return [found_char(char, options[raw.casefold()], source, "pair", fragment)]
    item: Fragment = {
        "form": "words",
        "name": None,
        "value": raw,
        "text": raw,
        "number": None,
        "unit": None,
        "start": 0,
        "end": len(raw),
    }
    raw_words = words(raw)
    pair_options = [o for o in options_of([char], source) if o["value"] != "Да"]
    found = chosen(
        explain(raw_words, pair_options, prior, set()), item, raw_words, "pair", multiple
    )
    if found:
        for char_found in found:
            char_found["fragment"] = fragment
        return found
    left = found_char(char, raw, source, "pair", fragment)
    if source == "ktru":
        shown = ", ".join(sorted(set(options.values()))[:SHOWN_OPTIONS])
        mark(left, "not_in_list", options=shown)
    return [left]


def told_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    product: set[str],
    shared: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Значения без названия: «деревянный» → Тип каркаса: Деревянный. Значение портала берём,
    только если рядом названа и сама характеристика: в списках портала много мусора."""
    ktru = options_of(chars, "ktru")
    portal = options_of(portal_chars, "portal")
    multiple = multiple_names(chars)
    found = []
    for item in fragments:
        if item["form"] == "words":
            fragment_words = words(item["text"])
            found += chosen(
                explain(fragment_words, ktru, prior, shared),
                item,
                fragment_words,
                "value",
                multiple,
            )
            # Слова товара не подсказывают характеристику портала: «Вид стола: Письменный».
            told = [word for word in fragment_words if word["stem"] not in product]
            found += chosen(explain(told, portal, prior, shared), item, told, "value", multiple)
    return found


def searched_text(char: Characteristic, value: str) -> str:
    """«Да» ищем по названию: «из натуральной кожи» значит «Натуральная кожа: Да»."""
    if value == "Да":
        return char["name"]
    return value


def number_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
) -> list[FoundChar]:
    """Числа в единице КТРУ: «2,2 см» → Толщина столешницы 22 мм, «высота 75» → 75 см."""
    numeric = numeric_chars(chars)
    portal_numeric = numeric_chars(portal_chars)
    found: list[FoundChar] = []
    for item in fragments:
        if item["form"] != "number" or item["number"] is None:
            continue
        name_stems = stems(item["name"] or "")
        char = owner(numeric, item["unit"], name_stems, guess=True)
        source: Source = "ktru"
        if char is None:
            # В названиях портала много мусора, поэтому название должно быть в тексте целиком.
            told = name_stems - product
            named_here = [c for c in portal_numeric if stems(c["name"]) <= told]
            char = owner(named_here, item["unit"], told, guess=False)
            source = "portal"
        if char is None:
            continue
        values = allowed.get(char["name"], [])
        char_found = number_char(
            char, source, "number", item["number"], item["unit"], values, item["text"]
        )
        char_found["start"], char_found["end"] = item["start"], item["end"]
        found.append(char_found)
    return found


def numeric_chars(chars: list[Characteristic]) -> list[Characteristic]:
    numeric = {}
    for char in chars:
        if char["type"] == "number" and "unit" in char:
            numeric[char["name"]] = char
    return list(numeric.values())


def number_char(
    char: Characteristic,
    source: Source,
    rule: Rule,
    number: float,
    code: str | None,
    allowed: list[Value],
    fragment: str,
) -> FoundChar:
    unit: Unit | None = char["unit"] if "unit" in char else None
    value = number
    if unit and code:
        value = convert(number, code, unit["code"])
    item = found_char(char, tidy(value), source, rule, fragment, unit=unit)
    if unit and code and code != unit["code"]:
        mark(item, "unit")
    check_range(item, allowed)
    return item


def check_range(item: FoundChar, allowed: list[Value]) -> None:
    """Число вне коридора КТРУ: правим, если сдвиг на порядок его объясняет, иначе помечаем."""
    corridor = bounds(allowed)
    value = item["value"]
    if corridor is None or isinstance(value, str) or within(value, corridor):
        return
    for shift in SHIFTS:
        if within(value * shift, corridor):
            item["value"] = tidy(value * shift)
            mark(item, "wrong_unit")
            return
    low, high = corridor
    mark(item, "out_of_range", low=show(low), high=show(high))


def within(value: float, corridor: tuple[float, float]) -> bool:
    return corridor[0] <= value <= corridor[1]


def owner(
    chars: list[Characteristic], code: str | None, name_stems: set[str], guess: bool
) -> Characteristic | None:
    """Характеристика той же величины, чьё название ближе к словам перед числом. Угадывать
    без подсказки в тексте можно, только если выбора нет, единица названа и guess разрешён."""
    best = None
    best_hits = 0
    same_dimension = 0
    for char in chars:
        if code is not None and scale(char["unit"]["code"])[0] != scale(code)[0]:
            continue
        same_dimension += 1
        hits = len(stems(char["name"]) & name_stems)
        if best is None or hits > best_hits:
            best = char
            best_hits = hits
    if best_hits == 0 and (not guess or same_dimension > 1 or code is None):
        return None
    return best
