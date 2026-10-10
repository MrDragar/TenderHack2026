import argparse
import html
import json
import os
import sys
import tempfile
from collections import Counter
from collections.abc import Iterable, Iterator, Sequence
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import TypedDict

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from contracts.ref import Portal, PortalChar

SOURCE_FILES = ("Одежда.xlsx", "Мебель.xlsx", "Инструменты.xlsx")
EXPECTED_HEADERS = {
    3: "Название характеристики",
    5: "Тип характеристики",
    6: "Значения характеристики",
    7: "Единица измерения",
}
COLUMN_COUNT = 8
VALUES_COLUMN = 6
MISSING_VALUE_PREFIX = "[не найдено значение id:"


class CharStats(TypedDict):
    rows: int
    names: Counter[str]
    types: Counter[str]
    units: Counter[str]
    values: dict[str, Counter[str]]


def normalize_text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def clean_cell(value: object) -> str:
    text = str(value or "").split("\t", 1)[0].split("_x000D_", 1)[0]
    return normalize_text(text)


def formula_value(value: str | ArrayFormula) -> str | None:
    formula = value.text if isinstance(value, ArrayFormula) else value
    if not isinstance(formula, str):
        return None
    text = formula.removeprefix("=").strip()
    if "_FV(" in text:
        return None
    quoted = text.removeprefix("+").strip()
    if len(quoted) >= 2 and quoted[0] == quoted[-1] == '"':
        return quoted[1:-1]
    return text


def clock_value(value: time | timedelta) -> str:
    if isinstance(value, time):
        hours, minutes, seconds = value.hour, value.minute, value.second
    else:
        total_seconds = int(value.total_seconds())
        hours, remainder = divmod(total_seconds, 3600)
        minutes, seconds = divmod(remainder, 60)
    result = f"{hours}:{minutes}"
    return f"{result}:{seconds}" if seconds else result


def portal_value(cell) -> object:
    value = cell.value
    if cell.data_type == "f":
        return formula_value(value)
    if isinstance(value, datetime | date):
        return None
    if isinstance(value, time | timedelta):
        return clock_value(value)
    return value


def normalize_name(value: object) -> str:
    name = clean_cell(value)
    parts = [part.strip() for part in name.split("|") if part.strip()]
    if parts and len({part.casefold() for part in parts}) == 1:
        return parts[0]
    return name


def is_continuation(row: Sequence[object]) -> bool:
    return bool(normalize_text(row[0])) and not normalize_text(row[1])


def stitched_rows(rows: Iterable[Sequence[object]]) -> Iterator[list[object]]:
    current: list[object] | None = None
    for source_row in rows:
        row = list(source_row[:COLUMN_COUNT])
        row.extend([None] * (COLUMN_COUNT - len(row)))
        if not any(normalize_text(value) for value in row):
            if current is not None:
                yield current
                current = None
            continue
        if is_continuation(row):
            if current is None:
                raise ValueError("Хвост значения найден до основной строки")
            current[VALUES_COLUMN] = f"{current[VALUES_COLUMN] or ''}{row[0]}"
            continue
        if current is not None:
            yield current
        current = row
    if current is not None:
        yield current


def workbook_rows(path: Path) -> Iterator[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError(f"В {path.name} нет второго листа")
        sheet = workbook.worksheets[1]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        for index, expected in EXPECTED_HEADERS.items():
            actual = normalize_text(header[index]) if len(header) > index else ""
            if actual != expected:
                raise ValueError(
                    f"В {path.name} ожидалась колонка {expected!r}, найдена {actual!r}"
                )
        rows = (
            [
                portal_value(cell) if index == VALUES_COLUMN else cell.value
                for index, cell in enumerate(cells)
            ]
            for cells in sheet.iter_rows(min_row=2, max_col=COLUMN_COUNT)
        )
        yield from stitched_rows(rows)
    finally:
        workbook.close()


def most_common(counter: Counter[str]) -> str:
    return sorted(
        counter,
        key=lambda value: (-counter[value], value.casefold(), value),
    )[0]


def clean_values(value: object) -> Iterator[str]:
    for part in html.unescape(clean_cell(value)).split(";"):
        normalized = normalize_text(part).removeprefix("=").strip()
        if normalized and MISSING_VALUE_PREFIX not in normalized.casefold():
            yield normalized


def aggregate(rows: Iterable[Sequence[object]]) -> Portal:
    grouped: dict[str, dict[str, CharStats]] = {}
    for row in rows:
        category = clean_cell(row[0])
        name = normalize_name(row[3])
        if not category or not name:
            continue
        category_stats = grouped.setdefault(category, {})
        stats = category_stats.setdefault(
            name.casefold(),
            {
                "rows": 0,
                "names": Counter(),
                "types": Counter(),
                "units": Counter(),
                "values": {},
            },
        )
        stats["rows"] += 1
        stats["names"][name] += 1
        characteristic_type = clean_cell(row[5])
        if characteristic_type:
            stats["types"][characteristic_type] += 1
        unit = clean_cell(row[7])
        if unit:
            stats["units"][unit] += 1
        for value in clean_values(row[VALUES_COLUMN]):
            stats["values"].setdefault(value.casefold(), Counter())[value] += 1

    portal: Portal = {}
    for category in sorted(grouped, key=str.casefold):
        ranked: list[tuple[int, PortalChar]] = []
        for stats in grouped[category].values():
            if not stats["types"]:
                continue
            values = [
                (sum(spellings.values()), most_common(spellings))
                for spellings in stats["values"].values()
            ]
            values.sort(key=lambda item: (-item[0], item[1].casefold(), item[1]))
            characteristic: PortalChar = {
                "name": most_common(stats["names"]),
                "type": most_common(stats["types"]),
                "unit": most_common(stats["units"]) if stats["units"] else None,
                "values": [value for _, value in values],
            }
            ranked.append((stats["rows"], characteristic))

        ranked.sort(key=lambda item: (-item[0], item[1]["name"].casefold(), item[1]["name"]))
        portal[category] = [characteristic for _, characteristic in ranked]
    return portal


def write_portal(portal: Portal, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=output_path.parent,
        prefix=f".{output_path.stem}.",
        suffix=".json",
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with temporary_path.open("w", encoding="utf-8") as output:
            json.dump(portal, output, ensure_ascii=False, separators=(",", ":"))
            output.write("\n")
        os.replace(temporary_path, output_path)
        output_path.chmod(0o644)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise


def build(input_dir: Path, output_path: Path) -> Portal:
    missing = [filename for filename in SOURCE_FILES if not (input_dir / filename).is_file()]
    if missing:
        raise FileNotFoundError(f"Не найдены файлы: {', '.join(missing)}")

    portal = aggregate(
        row for filename in SOURCE_FILES for row in workbook_rows(input_dir / filename)
    )
    write_portal(portal, output_path)
    return portal


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Собирает характеристики портала со вторых листов трех XLSX."
    )
    parser.add_argument("--input", type=Path, default=Path("raw_data"))
    parser.add_argument("--output", type=Path, default=Path("data/processed/portal.json"))
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        portal = build(args.input, args.output)
    except (OSError, TypeError, ValueError) as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        return 1
    summary = {
        "output": str(args.output),
        "categories": len(portal),
        "characteristics": sum(map(len, portal.values())),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
