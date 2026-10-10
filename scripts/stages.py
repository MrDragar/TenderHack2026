import sys
from pathlib import Path
from random import Random

from openpyxl import load_workbook

from contracts.pipeline import Candidate, FoundChar, Fragment
from contracts.ref import Portal, Position
from pipeline.candidates import find
from pipeline.extract import NUMBER, extract
from pipeline.link import link
from pipeline.narrow import narrow
from pipeline.section import find_section
from pipeline.values import show
from scripts.generate import Generated, generate, gold_char, inside, section_texts, siblings

RAW_DIR = Path("data/raw/Казань 2026/Примеры выгрузок")
SHEETS = ("Одежда.xlsx", "Мебель.xlsx", "Инструменты.xlsx")
PRODUCTS_PER_SHEET = 200
GENERATED = 300
SECTIONS = 300

type Pairs = list[tuple[str, str]]


def share(part: int, total: int) -> str:
    return f"{part}/{total} ({part / total:.0%})" if total else "0/0"


def stage1(generated: list[Generated], ref: list[Position]) -> str:
    """Соседние коды одного шаблона называются одинаково, их различает сужение, а не этап 1.
    Поэтому ранг считаем по шаблонам."""
    by_code = {p["code"]: p for p in ref}
    lines = []
    for label, part in split(generated):
        top1 = 0
        top5 = 0
        reciprocal = 0.0
        category = 0
        sizes = 0
        for example in part:
            candidates = find(example["text"], ref)
            sizes += len(candidates)
            place = template_rank(example["code"], candidates, by_code)
            if place is not None:
                top1 += place == 1
                top5 += place <= 5
                reciprocal += 1 / place
            gold = by_code[example["code"]]["category"]
            category += bool(candidates) and by_code[candidates[0]["code"]]["category"] == gold
        total = len(part)
        mrr = reciprocal / total if total else 0
        mean = sizes / total if total else 0
        lines.append(
            f"{label}: шаблон в топ-1 {share(top1, total)}, в топ-5 {share(top5, total)}, "
            f"MRR {mrr:.2f}, категория {share(category, total)}, кандидатов {mean:.1f}"
        )
    return "1. Позиция КТРУ\n" + "\n".join(f"     {line}" for line in lines)


def template_rank(
    code: str, candidates: list[Candidate], by_code: dict[str, Position]
) -> int | None:
    templates: list[str] = []
    for candidate in candidates:
        template = by_code[candidate["code"]]["parent_code"] or candidate["code"]
        if template not in templates:
            templates.append(template)
    gold = by_code[code]["parent_code"] or code
    if gold not in templates:
        return None
    return templates.index(gold) + 1


def split(generated: list[Generated]) -> list[tuple[str, list[Generated]]]:
    """Шум прячет регресс на чистых текстах, поэтому считаем их отдельно."""
    clean = [e for e in generated if not e["noisy"]]
    spoiled = [e for e in generated if e["noisy"]]
    return [("чистые", clean), ("с порчей", spoiled)]


def stage2(texts: list[tuple[str, str]], portal: Portal) -> str:
    """Отказ считаем отдельно от ошибки: «не знаю» лучше уверенно неверного раздела."""
    right = 0
    refused = 0
    for text, section in texts:
        found = find_section(text, portal)
        right += found == section
        refused += found is None
    answered = len(texts) - refused
    return (
        f"2. Раздел портала: угадан {share(right, len(texts))}, "
        f"отказ {share(refused, len(texts))}, верно среди ответивших {share(right, answered)}"
    )


def products(names: set[str], raw_dir: Path = RAW_DIR) -> list[tuple[str, Pairs]]:
    """Товары листа 1: название и характеристики через « | ». Между названием характеристики
    и значением разделителя нет, делим по самому длинному названию из листа 2."""
    found = []
    for sheet_name in SHEETS:
        book = load_workbook(raw_dir / sheet_name, read_only=True)
        taken = 0
        for row in book.worksheets[0].iter_rows(min_row=2, values_only=True):
            if taken == PRODUCTS_PER_SHEET:
                break
            pairs = split_pairs(str(row[5] or ""), names)
            if row[1] and pairs:
                found.append((str(row[1]), pairs))
                taken += 1
        book.close()
    return found


def split_pairs(line: str, names: set[str]) -> Pairs:
    pairs = []
    for part in line.split(" | "):
        name = longest_prefix(part, names)
        if name:
            pairs.append((part[: len(name)], part[len(name) :].strip()))
    return pairs


def longest_prefix(part: str, names: set[str]) -> str | None:
    words = part.split()
    for size in range(len(words) - 1, 0, -1):
        prefix = " ".join(words[:size])
        if prefix.casefold() in names:
            return prefix
    return None


def stage3(items: list[tuple[str, Pairs]]) -> str:
    """Пара найдена, если есть кусок с тем же названием и значением; числа сравниваем числом.
    Мягко: значение нашлось, а название нет, этап 4 ещё может его привязать."""
    expected = 0
    hits = 0
    soft = 0
    named = 0
    for name, pairs in items:
        text = ", ".join([name] + [f"{n} {v}" for n, v in pairs])
        fragments = extract(text)
        with_name = [f for f in fragments if f["name"]]
        named += len(with_name)
        for pair_name, value in pairs:
            expected += 1
            if any(same_pair(f, pair_name, value) for f in with_name):
                hits += 1
            elif any(same_value(f, value) for f in fragments):
                soft += 1
    return (
        f"3. Извлечение: полнота {share(hits, expected)}, точность {share(hits, named)}, "
        f"F1 {f_score(hits, named, expected, 1):.2f}, "
        f"ещё только значение без названия {share(soft, expected)}"
    )


def same_pair(item: Fragment, name: str, value: str) -> bool:
    if (item["name"] or "").casefold() != name.casefold():
        return False
    return same_value(item, value)


def same_value(item: Fragment, value: str) -> bool:
    number = NUMBER.search(value)
    if item["number"] is not None and number:
        return item["number"] == float(number["number"].replace(",", "."))
    return item["value"].casefold() == value.casefold()


def f_score(right: int, found: int, expected: int, beta: float) -> float:
    """beta 0.5 для правок: ложная правка чистого значения хуже пропущенной."""
    if not right:
        return 0.0
    precision = right / found
    recall = right / expected
    return (1 + beta**2) * precision * recall / (beta**2 * precision + recall)


def stage4(generated: list[Generated], by_code: dict[str, Position], ref: list[Position]) -> str:
    expected = 0
    right = 0
    right_unit = 0
    extra = 0
    corrupted = 0
    fixes = 0
    right_fixes = 0
    false_alarms = 0
    clean = 0
    by_kind: dict[str, list[int]] = {}
    for example in generated:
        position = by_code[example["code"]]
        candidates = siblings(position, by_code)
        found = link(example["fragments"], candidates, None, ref, {})
        names = {char["name"] for char in example["chars"]}
        extra += sum(char["name"] not in names for char in found)
        for char in example["chars"]:
            expected += 1
            got = [
                f for f in found if f["name"] == char["name"] and same(f["value"], char["value"])
            ]
            right += bool(got)
            right_unit += any(unit_ok(f, position) for f in got)
            fixed = [f for f in found if f["name"] == char["name"] and f["kind"]]
            fixes += bool(fixed)
            if "kind" not in char:
                clean += 1
                false_alarms += bool(fixed)
                continue
            corrupted += 1
            hit = any(f["kind"] == char["kind"] for f in got)
            right_fixes += hit
            counts = by_kind.setdefault(char["kind"], [0, 0])
            counts[0] += hit
            counts[1] += 1
    kinds = ", ".join(f"{kind} {share(hit, total)}" for kind, (hit, total) in by_kind.items())
    return (
        f"4. Привязка и унификация: название и значение {share(right, expected)}, "
        f"с единицей {share(right_unit, expected)}, "
        f"лишних {extra}\n"
        f"     правки: исправлено {share(right_fixes, corrupted)}, "
        f"верных среди правок {share(right_fixes, fixes)}, "
        f"F0.5 {f_score(right_fixes, fixes, corrupted, 0.5):.2f}, "
        f"ложных на чистых {share(false_alarms, clean)}\n"
        f"     по видам порчи: {kinds}"
    )


def unit_ok(found: FoundChar, position: Position) -> bool:
    """Единица КТРУ у числа, у текстового значения единицы нет."""
    for char in position["characteristics"]:
        if char["name"] == found["name"] and "unit" in char:
            return found["unit"] is not None and found["unit"]["code"] == char["unit"]["code"]
    return found["unit"] is None


def same(a: str | int | float, b: str | int | float) -> bool:
    return show(a).strip().casefold() == show(b).strip().casefold()


def stage5(generated: list[Generated], by_code: dict[str, Position], ref: list[Position]) -> str:
    """Со всеми неизменяемыми значениями код должен найтись. Без одного допустим вопрос,
    но не уверенный неверный код; после ответа на вопрос код должен найтись."""
    rng = Random(0)
    full_right = 0
    right = 0
    asked = 0
    silent = 0
    answered_right = 0
    required = 0
    required_shown = 0
    for example in generated:
        position = by_code[example["code"]]
        candidates = siblings(position, by_code)
        chars = fixed_chars(position)
        full_right += narrow(candidates, chars, ref)["code"] == example["code"]
        if chars:
            removed = chars.pop(rng.randrange(len(chars)))
            if is_required(removed["name"], position):
                required += 1
                missing = narrow(candidates, chars, ref)["missing_required"]
                required_shown += removed["name"] in missing
        result = narrow(candidates, chars, ref)
        right += result["code"] == example["code"]
        silent += result["code"] is not None and result["code"] != example["code"]
        if result["question"]:
            asked += 1
            answer = answer_char(result["question"]["char"], position)
            if answer:
                chars.append(answer)
            answered_right += narrow(candidates, chars, ref)["code"] == example["code"]
    total = len(generated)
    return (
        f"5. Сужение: со всеми значениями код верен {share(full_right, total)}\n"
        f"     без одного: код верен {share(right, total)}, вопрос {share(asked, total)}, "
        f"молча неверно {silent}, после ответа код верен {share(answered_right, asked)}\n"
        f"     убранная обязательная показана в недостающих {share(required_shown, required)}"
    )


def is_required(name: str, position: Position) -> bool:
    return any(c["name"] == name and c["required"] for c in position["characteristics"])


def answer_char(name: str, position: Position) -> FoundChar | None:
    for char in fixed_chars(position):
        if char["name"] == name:
            return char
    return None


def fixed_chars(position: Position) -> list[FoundChar]:
    found = []
    for char in position["characteristics"]:
        if char["selection"] != "fixed" or not char["values"]:
            continue
        first = char["values"][0]
        value = first if isinstance(first, str) else inside(first)
        if value is not None:
            found.append(gold_char({"name": char["name"], "value": value}))
    return found


def report(ref: list[Position], portal: Portal, with_products: bool = True) -> str:
    by_code = {p["code"]: p for p in ref}
    generated = generate(ref, GENERATED)
    lines = [
        "Этапы на правильном входе",
        f"  {stage1(generated, ref)}",
        f"  {stage2(section_texts(portal, SECTIONS), portal)}",
    ]
    if with_products and portal and RAW_DIR.exists():
        names = set()
        for chars in portal.values():
            for char in chars:
                names.add(char["name"].casefold())
        lines.append(f"  {stage3(products(names))}")
    lines += [f"  {stage4(generated, by_code, ref)}", f"  {stage5(generated, by_code, ref)}"]
    return "\n".join(lines)


if __name__ == "__main__":
    from pipeline.run import portal, reference

    print(report(reference(), portal(), with_products="--fast" not in sys.argv))
