"""Товары листа 1 с ответом без ручной разметки: пары « | » для извлечения, раздел
листа 2 по совпадению характеристик. Выборка собирается один раз в data/processed.

  uv run python -m scripts.products
"""

import json
import re
import sys
import zlib
from pathlib import Path
from typing import TypedDict

from openpyxl import load_workbook

from contracts.ref import Category, Portal
from scripts.generate import held_out

RAW_DIR = Path("data/raw/Казань 2026/Примеры выгрузок")
PRODUCTS_PATH = Path("data/processed/products.json")
SHEETS: dict[str, Category] = {
    "Одежда.xlsx": "clothing",
    "Мебель.xlsx": "furniture",
    "Инструменты.xlsx": "tools",
}
# Каждый сотый товар по хешу ID: около тысячи на категорию вразброс по всему листу,
# а не первые строки, где подряд идут одни рюкзаки.
SAMPLE = 100
# Четверть названий листа 1 длиннее девяти слов.
LONG_NAME = 10
# Латиница, артикул «КБС-011Т», бренд капслоком «ПРАКТИК», серия в кавычках.
MODEL = re.compile(r"[a-zA-Z]|[А-ЯЁ]+-?\d|[А-ЯЁ]{5,}|[\"«]")

type Pairs = list[tuple[str, str]]


class Product(TypedDict):
    id: str
    category: Category
    name: str
    pairs: Pairs
    section: str | None
    tags: list[str]
    test: bool


def products(portal: Portal, test: bool = False, build: bool = True) -> list[Product]:
    """Пусто, если выборки нет и собирать её не просили или нет листа 1."""
    if not PRODUCTS_PATH.exists():
        if not build or not RAW_DIR.exists() or not portal:
            return []
        save(sample(portal))
    found = json.loads(PRODUCTS_PATH.read_text(encoding="utf-8"))
    return [p for p in found if p["test"] == test]


def sample(portal: Portal, raw_dir: Path = RAW_DIR) -> list[Product]:
    found = []
    for sheet_name, category in SHEETS.items():
        book = load_workbook(raw_dir / sheet_name, read_only=True)
        goods, chars = book.worksheets
        sections = section_names(chars.iter_rows(min_row=2, values_only=True), portal)
        names = set().union(*sections.values()) if sections else set()
        for row in goods.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1] or zlib.crc32(str(row[0]).encode()) % SAMPLE:
                continue
            pairs = split_pairs(str(row[5] or ""), names)
            if pairs:
                found.append(product(row, category, pairs, sections))
        book.close()
    return found


def section_names(rows, portal: Portal) -> dict[str, set[str]]:
    """Названия как в листе 2 до унификации: так их пишут и в листе 1. Только разделы
    справочника, остальное это строки, разъехавшиеся при выгрузке."""
    sections: dict[str, set[str]] = {}
    for row in rows:
        if row[0] in portal and row[3]:
            sections.setdefault(row[0], set()).add(" ".join(str(row[3]).split()).casefold())
    return sections


def product(row: tuple, category: Category, pairs: Pairs, sections: dict[str, set[str]]) -> Product:
    name = " ".join(str(row[1]).split())
    return {
        "id": str(row[0]),
        "category": category,
        "name": name,
        "pairs": pairs,
        "section": weak_section(pairs, sections),
        "tags": tags(name),
        # Свой ключ: в выборку берём crc32 % 100 == 0, а с ним остаток от 5 всегда 0.
        "test": held_out(f"товар {row[0]}"),
    }


def weak_section(pairs: Pairs, sections: dict[str, set[str]]) -> str | None:
    """Товару в портале дают характеристики его раздела, поэтому раздел, где есть все
    названия товара, и есть его раздел. Таких два и больше или ни одного: метки нет,
    неточная метка хуже пропущенной."""
    told = {name.casefold() for name, _ in pairs}
    covering = [section for section, names in sections.items() if told <= names]
    return covering[0] if len(covering) == 1 else None


def tags(name: str) -> list[str]:
    found = []
    if len(name.split()) >= LONG_NAME:
        found.append("длинное название")
    if MODEL.search(name):
        found.append("бренд или модель")
    return found


def split_pairs(line: str, names: set[str]) -> Pairs:
    """Между названием характеристики и значением разделителя нет, делим по самому
    длинному известному названию."""
    pairs = []
    for part in line.split(" | "):
        name = longest_prefix(part, names)
        if name:
            pairs.append((name, " ".join(part.split()[len(name.split()) :])))
    return pairs


def longest_prefix(part: str, names: set[str]) -> str | None:
    words = part.split()
    for size in range(len(words) - 1, 0, -1):
        prefix = " ".join(words[:size])
        if prefix.casefold() in names:
            return prefix
    return None


def save(found: list[Product]) -> None:
    PRODUCTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    PRODUCTS_PATH.write_text(json.dumps(found, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    from pipeline.run import portal

    if not RAW_DIR.exists():
        sys.exit(f"Нет листа 1 в {RAW_DIR}")
    built = sample(portal())
    save(built)
    labeled = sum(p["section"] is not None for p in built)
    print(f"Товаров {len(built)}, с разделом {labeled}, в {PRODUCTS_PATH}")
