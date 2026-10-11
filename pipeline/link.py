"""Значение характеристики это то, к которому кусок текста приводится самой дешёвой порчей.

Зашумлённый канал (Brill и Moore 2000): цена значения это сумма по его словам самой дешёвой
порчи (чисто, раскладка, форма слова, правка букв) плюс пропуск слов значения и выбор среди
значений кандидатов. Цены порч это частоты на листах 2: написание базовой характеристики
против значения связанной эталонной. Слово, которое значение не объясняет, стоит как шум
из словаря поставщиков листа 1. Отрицание (NegEx, Chapman 2001) тоже порча: «без подсветки»
дёшево объясняется «Наличие подсветки: Нет», окно и цены посчитаны по листу 1. Значение
отдаём, только если P ≥ 1/2. Числа привязываются по словам перед числом и единице.
"""

import json
import re
import zlib
from collections import Counter
from collections.abc import Callable
from difflib import SequenceMatcher
from functools import cache
from math import exp, inf, log
from pathlib import Path
from typing import Literal, TypedDict

from openpyxl import load_workbook

from contracts.pipeline import Candidate, FixKind, FoundChar, Fragment, Rule
from contracts.ref import Characteristic, Portal, PortalChar, Position, Unit, Value
from pipeline.data import reference
from pipeline.explain import mark
from pipeline.text import Word, stem, stems, words
from pipeline.values import UNIT_ALIASES, bounds, convert, same, scale, show, tidy

# Перепутанная единица чаще всего сдвигает число на порядок: см вместо мм, м вместо см.
SHIFTS = (10, 0.1, 100, 0.01, 1000, 0.001)
SHOWN_OPTIONS = 6
COSTS_PATH = Path("data/processed/link_noisy_channel_costs.json")
# Написание базовой характеристики сравниваем с эталонным, только если это одно значение.
SAME_VALUE = 0.5
MAX_EDITS = 3
RUSSIAN = "абвгдежзийклмнопрстуфхцчшщъыьэюя"
ALPHABET = len(RUSSIAN) + 1
OPS = ("sub", "del", "ins", "swap")
NEGATION_PATH = Path("data/processed/link_negation_counts.json")
# Грамматические слова отрицания, действуют вперёд (NegEx): «без подсветки», «не складной».
TRIGGERS = {"без", "не", "нет", "отсутствует", "отсутствуют", "отсутствие"}
MAX_SCOPE = 6
# Знаки препинания обрывают действие отрицания, как конец фразы в NegEx.
TOKEN = re.compile(r"\w+|[,;:.()]")

UNIT_IN_NAME = re.compile(r"[(,]\s*(?P<unit>[^(),]+?)\)?$")

type Source = Literal["ktru", "portal"]


class Option(TypedDict):
    char: Characteristic
    value: str
    source: Source


class Explained(TypedDict):
    option: Option
    gain: float
    used: set[int]
    named: set[int]
    kind: FixKind | None


def link(
    fragments: list[Fragment],
    candidates: list[Candidate],
    section: str | None,
    ref: list[Position],
    portal: Portal,
) -> list[FoundChar]:
    by_code = {p["code"]: p for p in ref}
    positions = [by_code[candidate["code"]] for candidate in candidates]
    chars = [char for position in positions for char in position["characteristics"]]
    allowed = allowed_values(chars)
    portal_chars = [as_ktru(char) for char in portal.get(section or "", [])]
    product = stems(" ".join(c["name"] for c in candidates))
    prior = value_prior(positions)

    pairs = told_pairs(fragments, chars, portal_chars, allowed, product, prior)
    shared = set.intersection(*(stems(c["name"]) for c in candidates)) if candidates else set()
    values = told_values(fragments, chars, portal_chars, product, shared, prior)
    numbers = number_values(fragments, chars, portal_chars, allowed, product)
    found = choose(pairs, values, numbers)
    # Обязательность у соседних кодов разная: обязательна, если обязательна хоть у одного.
    required = {char["name"] for char in chars if char["required"]}
    for char in found:
        char["required"] = char["source"] == "ktru" and char["name"] in required
    return found


def choose(
    pairs: list[FoundChar], values: list[FoundChar], numbers: list[FoundChar]
) -> list[FoundChar]:
    """Явная пара сильнее значения без названия, оно сильнее числа. Значения одного правила
    уже выбраны каналом, между правилами разные значения дают правку «два значения»."""
    by_name: dict[str, list[FoundChar]] = {}
    for rule in (pairs, values, numbers):
        mine: dict[str, list[FoundChar]] = {}
        for char in rule:
            mine.setdefault(char["name"], []).append(char)
        for name, chars in mine.items():
            kept = by_name.get(name)
            if kept is None:
                by_name[name] = chars
                continue
            other = [c for c in chars if not any(same(k["value"], c["value"]) for k in kept)]
            if other:
                mark(kept[0], "conflict", other=other[0]["fragment"])
    return [char for chars in by_name.values() for char in chars]


@cache
def costs() -> dict[str, float]:
    """−log частоты каждой порчи. Счётчики собираются один раз по листам 2 в data/processed."""
    if not COSTS_PATH.exists():
        COSTS_PATH.write_text(json.dumps(mined_counts(), ensure_ascii=False, indent=1))
    counts = json.loads(COSTS_PATH.read_text())
    return priced(Counter(counts["words"]), Counter(counts["edits"]))


@cache
def noise() -> float:
    """Слово шума это порча «лишнее слово» и любое слово из тех, что поставщики пишут в товарах
    листа 1, на выбор."""
    # Лист 1 читаем здесь, а не при импорте: scripts.products тянет генератор, а он этот модуль.
    from pipeline.data import portal
    from scripts.products import products

    vocabulary = set()
    for product in products(portal()):
        vocabulary |= {word["text"] for word in words(product["name"])}
        for name, value in product["pairs"]:
            vocabulary |= {word["text"] for word in words(f"{name} {value}")}
    # Без листа 1 выборку не собрать, считаем словарь из одного слова, а не падаем.
    return costs()["noise"] + log(max(len(vocabulary), 1))


@cache
def negation() -> dict[str, float]:
    """Цены порчи «отрицание»: −log P(отрицание в тексте | ответ Да или Нет), −log P(ответ)
    и окно, на котором отрицание чаще значит «Нет», чем «Да». Всё из листа 1: название товара
    против его же пары «Наличие X: Да/Нет»."""
    if not NEGATION_PATH.exists():
        NEGATION_PATH.write_text(json.dumps(negation_counts(), indent=1))
    counts = Counter(json.loads(NEGATION_PATH.read_text()))
    window = 1
    for distance in range(1, MAX_SCOPE + 1):
        if counts[f"Нет:{distance}"] > counts[f"Да:{distance}"]:
            window = distance
    found = {"window": float(window)}
    totals = {}
    for answer in ("Да", "Нет"):
        near = sum(counts[f"{answer}:{d}"] for d in range(1, window + 1))
        total = near + counts[f"{answer}:0"]
        found[f"{answer}:True"] = -log((near + 1) / (total + 2))
        found[f"{answer}:False"] = -log((total - near + 1) / (total + 2))
        totals[answer] = total
    for answer, total in totals.items():
        found[answer] = -log(total / sum(totals.values()))
    return found


def negation_counts() -> Counter[str]:
    """«ответ:расстояние» от слова названия характеристики до слова отрицания перед ним,
    0 если отрицания нет. Товары выборки замера пропускаем."""
    from scripts.products import RAW_DIR, SAMPLE, SHEETS

    counts: Counter[str] = Counter()
    for sheet in SHEETS:
        book = load_workbook(RAW_DIR / sheet, read_only=True)
        for row in book.worksheets[0].iter_rows(min_row=2, max_col=6, values_only=True):
            if not row[1] or not row[5] or zlib.crc32(str(row[0]).encode()) % SAMPLE == 0:
                continue
            tokens = TOKEN.findall(str(row[1]).lower())
            for part in str(row[5]).split(" | "):
                said = part.split()
                if len(said) < 2 or said[-1].capitalize() not in ("Да", "Нет"):
                    continue
                distance = trigger_distance(tokens, stems(" ".join(said[:-1])) - stems("наличие"))
                if distance is not None:
                    counts[f"{said[-1].capitalize()}:{distance}"] += 1
        book.close()
    return counts


def trigger_distance(tokens: list[str], name_stems: set[str]) -> int | None:
    for i, token in enumerate(tokens):
        if len(token) >= 3 and stem(token) in name_stems:
            return scope_distance(tokens, i)
    return None


def scope_distance(tokens: list[str], index: int) -> int:
    """Сколько слов до ближайшего отрицания слева в пределах фразы, 0 если его нет."""
    for back in range(1, MAX_SCOPE + 1):
        if index - back < 0 or not tokens[index - back][0].isalnum():
            return 0
        if tokens[index - back] in TRIGGERS:
            return back
    return 0


def negated_words(text: str, fragment_words: list[Word]) -> set[int]:
    """Слова текста в окне действия отрицания."""
    window = int(negation()["window"])
    spans = [(m.start(), m[0].lower()) for m in TOKEN.finditer(text)]
    tokens = [token for _, token in spans]
    starts = [start for start, _ in spans]
    found = set()
    for i, word in enumerate(fragment_words):
        if word["start"] in starts:
            distance = scope_distance(tokens, starts.index(word["start"]))
            if 0 < distance <= window:
                found.add(i)
    return found


def says_negation(value: str) -> bool:
    return value == "Нет" or bool(TRIGGERS & set(TOKEN.findall(value.lower())))


def mined_counts() -> dict[str, Counter[str]]:
    from scripts.products import RAW_DIR, SHEETS

    written: Counter[tuple[str, tuple[str, ...]]] = Counter()
    for sheet in SHEETS:
        book = load_workbook(RAW_DIR / sheet, read_only=True)
        for row in book.worksheets[1].iter_rows(min_row=2, max_col=11, values_only=True):
            linked = tuple(linked_values(str(row[10] or "")))
            if row[1] and row[6] and linked:
                for value in split_values(str(row[6])):
                    written[value, linked] += 1
        book.close()
    words_seen: Counter[str] = Counter()
    edits: Counter[str] = Counter()
    for (value, linked), times in written.items():
        count_spoils(value, linked, times, words_seen, edits)
    return {"words": words_seen, "edits": edits}


def priced(words_seen: Counter[str], edits: Counter[str]) -> dict[str, float]:
    """Цена вида порчи по словам, цена правки по виду и по букве: P(буква текста | буква КТРУ)
    для замены и P(буква) для вставки, со сглаживанием на букву алфавита."""
    said = sum(words_seen[k] for k in ("clean", "layout", "form", "edit", "noise"))
    found = {}
    for kind in ("clean", "layout", "form", "edit", "noise"):
        found[kind] = -log((words_seen[kind] + 1) / said)
    found["omit"] = -log((words_seen["omit"] + 1) / (words_seen["kept"] + words_seen["omit"]))
    by_kind: Counter[str] = Counter()
    by_letter: Counter[str] = Counter()
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        by_kind[kind] += count
        if kind == "sub":
            by_letter["sub:" + letters[0]] += count
    total = sum(by_kind[kind] for kind in OPS) + len(OPS)
    for kind in OPS:
        found[kind] = -log((by_kind[kind] + 1) / total)
    for op, count in edits.items():
        kind, _, letters = op.partition(":")
        if kind == "sub":
            found[op] = -log((count + 1) / (by_letter["sub:" + letters[0]] + ALPHABET))
        elif kind == "ins":
            found[op] = -log((count + 1) / (by_kind["ins"] + ALPHABET))
        elif kind == "form":
            found[op] = 0.0
    # Пара окончаний, которой не было в листах 2, это всё ещё форма слова, а не опечатка:
    # её доля это доля впервые встреченных пар (Witten и Bell 1991).
    forms = [count for op, count in edits.items() if op.startswith("form:")]
    found["form:"] = -log(len(forms) / (sum(forms) + len(forms)))
    for letter in RUSSIAN:
        found.setdefault("sub:" + letter, -log(1 / (by_letter["sub:" + letter] + ALPHABET)))
    found["ins:"] = -log(1 / (by_kind["ins"] + ALPHABET))
    return found


def op_price(op: str) -> float:
    price = costs()
    kind, _, letters = op.partition(":")
    if kind == "sub":
        letter = price.get(op, price.get("sub:" + letters[0], price["ins:"]))
        return price[kind] + letter
    if kind == "ins":
        return price[kind] + price.get(op, price["ins:"])
    return price[kind]


def split_values(cell: str) -> list[str]:
    return [part.strip() for part in cell.split(";") if part.strip()]


def linked_values(cell: str) -> list[str]:
    """«5517160: М0; М1 | 5517205: М0; М1» → значения эталонных характеристик."""
    found = []
    for part in cell.split(" | "):
        found += split_values(part.split(":", 1)[-1])
    return found


def count_spoils(
    value: str, linked: tuple[str, ...], times: int, words_seen: Counter, edits: Counter
) -> None:
    folded = value.casefold()
    best = max(linked, key=lambda v: SequenceMatcher(None, folded, v.casefold()).ratio())
    if SequenceMatcher(None, folded, best.casefold()).ratio() < SAME_VALUE:
        return
    meant = [word["text"] for word in words(best)]
    hit = set()
    for word in words(value):
        kind, ops, index = None, [], -1
        for i, target in enumerate(meant):
            found, found_ops = spoil_kind(word, target)
            if found and (kind is None or RANK[found] < RANK[kind]):
                kind, ops, index = found, found_ops, i
        words_seen[kind or "noise"] += times
        for op in ops:
            edits[op] += times
        hit.add(index)
    kept = len(hit - {-1})
    words_seen["kept"] += kept * times
    words_seen["omit"] += (len(meant) - kept) * times


RANK = {"clean": 0, "layout": 1, "form": 2, "edit": 3}


def spoil_kind(word: Word, target: str) -> tuple[str | None, list[str]]:
    """Правки формы слова считаем как пару окончаний «ый>ая», правки строки по видам."""
    if word["text"] == target:
        return ("layout" if word["fix"] else "clean"), []
    if word["stem"] == stem(target):
        return "form", ["form:" + tails(word["text"], target)]
    distance, ops = edit(word["text"], target, unit)
    if distance <= 2 and len(target) >= 4:
        return "edit", ops
    return None, []


def same_word(said: str, meant: str) -> bool:
    """Падеж и род меняют до трёх последних букв: «чехлом» и «чехла», «стальной» и «сталь».
    Обрезка до шести букв в stem их разводит."""
    root = len(meant) - len(tails(said, meant).split(">")[0])
    return root >= 4 and len(meant) - root <= 3 and len(said) - root <= 3


def tails(said: str, meant: str) -> str:
    common = 0
    while common < min(len(said), len(meant)) and said[common] == meant[common]:
        common += 1
    return f"{meant[common:]}>{said[common:]}"


def unit(_: str) -> float:
    return 1.0


def edit(said: str, meant: str, price: Callable[[str], float]) -> tuple[float, list[str]]:
    """Цена правок, которыми meant превращается в said: замена, пропуск, вставка, перестановка
    соседних букв (Дамерау). Возвращает и сами правки, по ним считаются частоты."""
    if abs(len(said) - len(meant)) > MAX_EDITS:
        return inf, []
    rows, cols = len(meant) + 1, len(said) + 1
    table: list[list[tuple[float, tuple[str, ...]]]] = [[(0.0, ())] * cols for _ in range(rows)]
    for i in range(1, rows):
        table[i][0] = step(table[i - 1][0], f"del:{meant[i - 1]}", price)
    for j in range(1, cols):
        table[0][j] = step(table[0][j - 1], f"ins:{said[j - 1]}", price)
    for i in range(1, rows):
        for j in range(1, cols):
            steps = [
                step(table[i - 1][j], f"del:{meant[i - 1]}", price),
                step(table[i][j - 1], f"ins:{said[j - 1]}", price),
            ]
            if meant[i - 1] == said[j - 1]:
                steps.append(table[i - 1][j - 1])
            else:
                steps.append(step(table[i - 1][j - 1], f"sub:{meant[i - 1]}{said[j - 1]}", price))
            if i > 1 and j > 1 and meant[i - 1] == said[j - 2] and meant[i - 2] == said[j - 1]:
                steps.append(step(table[i - 2][j - 2], "swap", price))
            table[i][j] = min(steps, key=lambda found: found[0])
    cost, ops = table[-1][-1]
    return cost, list(ops)


def step(
    before: tuple[float, tuple[str, ...]], op: str, price: Callable[[str], float]
) -> tuple[float, tuple[str, ...]]:
    return before[0] + price(op), (*before[1], op)


@cache
def word_cost(said: str, layout: bool, meant: str) -> tuple[float, FixKind | None]:
    """Самая дешёвая порча, которой слово КТРУ meant стало словом текста said."""
    price = costs()
    ways: list[tuple[float, FixKind | None]] = [
        (price["edit"] + edit(said, meant, op_price)[0], "typo"),
        (price["form"] + price["edit"] + edit(stem(said), stem(meant), op_price)[0], "typo"),
    ]
    if said == meant:
        ways.append((price["clean"], None))
    if same_word(said, meant):
        ways.append((price["form"] + price.get("form:" + tails(said, meant), price["form:"]), None))
    cost, kind = min(ways, key=lambda way: way[0])
    if layout:
        return cost + price["layout"] - price["clean"], "layout"
    return cost, kind


def value_prior(positions: list[Position]) -> dict[tuple[str, str], float]:
    """−log доли кандидатов, у которых есть это значение, среди тех, у кого есть характеристика."""
    having: Counter[str] = Counter()
    valued: Counter[tuple[str, str]] = Counter()
    for position in positions:
        for char in position["characteristics"]:
            having[char["name"]] += 1
            for value in {v for v in char["values"] if isinstance(v, str)}:
                valued[char["name"], value] += 1
    return {key: -log(count / having[key[0]]) for key, count in valued.items()}


def explain(
    fragment_words: list[Word],
    options: list[Option],
    prior: dict[tuple[str, str], float],
    product: set[str],
    negated: set[int],
) -> list[Explained]:
    """Жадно объясняем текст: значение с наибольшим выигрышем над шумом забирает свои слова
    и слова названия своей характеристики, остальные пересчитываются на оставшихся. Слова,
    общие для названий всех кандидатов, уже объяснены товаром. Если соперник за те же слова
    почти так же вероятен (P лучшего < 1/2), слова забираем, а ответа не даём."""
    taken = {i for i, word in enumerate(fragment_words) if word["stem"] in product}
    values: set[int] = set()
    found = []
    while True:
        scored = []
        for option in options:
            match = score(fragment_words, option, taken, values, prior, log(len(options)), negated)
            if match:
                scored.append(match)
        if not scored:
            return found
        best = max(scored, key=lambda s: s["gain"])
        rivals = [s for s in scored if s is not best and s["used"] & best["used"]]
        odds = exp(-best["gain"]) + sum(exp(s["gain"] - best["gain"]) for s in rivals)
        taken |= best["used"] | best["named"]
        values |= best["used"]
        if 1 / (1 + odds) >= 0.5:
            found.append(best)


def score(
    fragment_words: list[Word],
    option: Option,
    taken: set[int],
    values: set[int],
    prior: dict[tuple[str, str], float],
    choice: float,
    negated: set[int],
) -> Explained | None:
    """Выигрыш значения над шумом: −log P(z | шум) − (−log P(z | a) − log P(a)), где P(a)
    это выбор одного из значений кандидатов и доля кандидатов с этим значением."""
    price = costs()
    meant = [word["text"] for word in words(searched_text(option["char"], option["value"]))]
    if not meant:
        return None
    gain = -choice - prior.get((option["char"]["name"], option["value"]), 0.0)
    used: set[int] = set()
    kinds = []
    for target in meant:
        place, cost, kind = closest(fragment_words, target, taken | used)
        if place is None:
            gain -= price["omit"]
            continue
        used.add(place)
        gain += noise() - cost
        kinds.append(kind)
    if not used:
        return None
    gain -= negation_cost(option["value"], bool(used & negated))
    named: set[int] = set()
    if option["value"] not in ("Да", "Нет"):
        for target in {word["text"] for word in words(option["char"]["name"])}:
            place, cost, _ = closest(fragment_words, target, values | used | named)
            if place is not None:
                named.add(place)
                gain += noise() - cost
        if option["source"] == "portal" and not named:
            return None
    if gain <= 0:
        return None
    kind = "layout" if "layout" in kinds else "typo" if "typo" in kinds else None
    return {"option": option, "gain": gain, "used": used, "named": named, "kind": kind}


def negation_cost(value: str, negated: bool) -> float:
    """Отрицание в тексте и в значении должны совпасть: «без подсветки» дёшево объясняется
    «Нет», дорого «Да». У Да и Нет ещё своя доля в листе 1."""
    price = negation()
    answer = "Нет" if says_negation(value) else "Да"
    cost = price[f"{answer}:{negated}"]
    if value in ("Да", "Нет"):
        cost += price[value]
    return cost


def closest(
    fragment_words: list[Word], target: str, busy: set[int]
) -> tuple[int | None, float, FixKind | None]:
    """Слово текста, которое дешевле шума объясняется словом target."""
    best: tuple[int | None, float, FixKind | None] = (None, noise(), None)
    for i, word in enumerate(fragment_words):
        if i in busy:
            continue
        cost, kind = word_cost(word["text"], word["fix"] == "layout", target)
        if cost < best[1]:
            best = (i, cost, kind)
    return best


def options_of(chars: list[Characteristic], source: Source) -> list[Option]:
    seen = set()
    found: list[Option] = []
    for char in chars:
        for value in char["values"]:
            if isinstance(value, str) and (char["name"], value) not in seen:
                if value == "Нет" and not all(v in ("Да", "Нет") for v in char["values"]):
                    continue
                seen.add((char["name"], value))
                found.append({"char": char, "value": value, "source": source})
    return found


def chosen(
    explained: list[Explained],
    item: Fragment,
    fragment_words: list[Word],
    rule: Rule,
    multiple: set[str],
) -> list[FoundChar]:
    """Одно значение на характеристику, несколько только там, где КТРУ разрешает выбор."""
    found: list[FoundChar] = []
    for match in explained:
        option = match["option"]
        name = option["char"]["name"]
        if any(f["name"] == name for f in found) and name not in multiple:
            continue
        start = min(fragment_words[i]["start"] for i in match["used"])
        end = max(fragment_words[i]["end"] for i in match["used"])
        char = found_char(
            option["char"],
            option["value"],
            option["source"],
            rule,
            item["text"][start:end],
            item["start"] + start,
            item["start"] + end,
        )
        if match["kind"]:
            mark(char, match["kind"])
        found.append(char)
    return found


def found_char(
    char: Characteristic,
    value: str | int | float,
    source: Source,
    rule: Rule,
    fragment: str,
    start: int = 0,
    end: int = 0,
    unit: Unit | None = None,
) -> FoundChar:
    return {
        "name": char["name"],
        "value": value,
        "unit": unit,
        "source": source,
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": char["required"],
        "fix": None,
        "kind": None,
        "rule": rule,
    }


def as_ktru(char: PortalChar) -> Characteristic:
    """Характеристика портала в виде КТРУ, чтобы искать и переводить единицы одним кодом."""
    numeric = "числов" in char["type"].lower()
    item: Characteristic = {
        "code": "",
        "name": char["name"],
        "type": "number" if numeric else "text",
        "selection": "single",
        "required": False,
        "values": [] if numeric else list(char["values"]),
    }
    unit = (char["unit"] or "").strip("() ")
    if unit:
        item["unit"] = {"code": UNIT_ALIASES.get(unit.lower(), unit), "name": unit}
    return item


def allowed_values(chars: list[Characteristic]) -> dict[str, list[Value]]:
    """Значения характеристики по всем кандидатам: у соседних кодов списки разные."""
    allowed: dict[str, list[Value]] = {}
    for char in chars:
        allowed.setdefault(char["name"], []).extend(char["values"])
    return allowed


def multiple_names(chars: list[Characteristic]) -> set[str]:
    return {char["name"] for char in chars if char["selection"] == "multiple"}


def told_pairs(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Человек дописывает исправление в конец, поэтому из явных значений верны последние."""
    last: dict[str, list[FoundChar]] = {}
    for item in fragments:
        if item["form"] == "pair":
            found = pair_chars(item, chars, portal_chars, allowed, product, prior)
            if found:
                last[found[0]["name"]] = found
    return [char for found in last.values() for char in found]


def pair_chars(
    item: Fragment,
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    offset = name_offset(item["name"] or "", product)
    if offset is None:
        return []
    name = (item["name"] or "")[offset:]
    start = item["start"] + offset
    fragment = item["text"][offset:]
    char = named(chars, name)
    source: Source = "ktru"
    if char is None:
        char = named(portal_chars, name)
        source = "portal"
    if char is None:
        return [new_char(name, item["value"], fragment, start, item["end"])]
    values = allowed.get(char["name"], []) if source == "ktru" else char["values"]
    if char["type"] == "number":
        if item["number"] is None:
            return []
        found = [number_char(char, source, "pair", item["number"], item["unit"], values, fragment)]
    else:
        whole = char.copy()
        whole["values"] = values
        found = text_pair(whole, source, item["value"], fragment, prior, multiple_names(chars))
    for char_found in found:
        char_found["start"], char_found["end"] = start, item["end"]
    return found


def name_offset(name: str, product: set[str]) -> int | None:
    """«Джемпер рост: 170»: слова товара перед названием характеристики к нему не относятся."""
    for word in words(name):
        if word["stem"] not in product:
            return word["start"]
    return None


def named(chars: list[Characteristic], name: str) -> Characteristic | None:
    """Характеристика с тем же названием, иначе та, в названии которой есть все слова."""
    for char in chars:
        if char["name"].casefold() == name.casefold():
            return char
    wanted = stems(name)
    for char in chars:
        if wanted and wanted <= stems(char["name"]):
            return char
    return None


def new_char(name: str, raw: str, fragment: str, start: int, end: int) -> FoundChar:
    return {
        "name": name[:1].upper() + name[1:],
        "value": raw,
        "unit": None,
        "source": "new",
        "fragment": fragment,
        "start": start,
        "end": end,
        "required": False,
        "fix": None,
        "kind": None,
        "rule": "pair",
    }


def text_pair(
    char: Characteristic,
    source: Source,
    raw: str,
    fragment: str,
    prior: dict[tuple[str, str], float],
    multiple: set[str],
) -> list[FoundChar]:
    """Значение пары это те значения списка, к которым текст приводится дешевле шума. Списки
    портала собраны из того, что вводили поставщики, поэтому «нет в списке» только про КТРУ."""
    options = {v.casefold(): v for v in char["values"] if isinstance(v, str)}
    if raw.casefold() in options:
        return [found_char(char, options[raw.casefold()], source, "pair", fragment)]
    item: Fragment = {
        "form": "words",
        "name": None,
        "value": raw,
        "text": raw,
        "number": None,
        "unit": None,
        "start": 0,
        "end": len(raw),
    }
    raw_words = words(raw)
    pair_options = [o for o in options_of([char], source) if o["value"] not in ("Да", "Нет")]
    negated = negated_words(raw, raw_words)
    found = chosen(
        explain(raw_words, pair_options, prior, set(), negated), item, raw_words, "pair", multiple
    )
    if found:
        for char_found in found:
            char_found["fragment"] = fragment
        return found
    left = found_char(char, raw, source, "pair", fragment)
    if source == "ktru":
        shown = ", ".join(sorted(set(options.values()))[:SHOWN_OPTIONS])
        mark(left, "not_in_list", options=shown)
    return [left]


def told_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    product: set[str],
    shared: set[str],
    prior: dict[tuple[str, str], float],
) -> list[FoundChar]:
    """Значения без названия: «деревянный» → Тип каркаса: Деревянный. Значение портала берём,
    только если рядом названа и сама характеристика: в списках портала много мусора."""
    ktru = options_of(chars, "ktru")
    portal = options_of(portal_chars, "portal")
    multiple = multiple_names(chars)
    found = []
    for item in fragments:
        if item["form"] == "words":
            fragment_words = words(item["text"])
            negated = negated_words(item["text"], fragment_words)
            found += chosen(
                explain(fragment_words, ktru, prior, shared, negated),
                item,
                fragment_words,
                "value",
                multiple,
            )
            # Слова товара не подсказывают характеристику портала: «Вид стола: Письменный».
            told = [word for word in fragment_words if word["stem"] not in product]
            told_negated = negated_words(item["text"], told)
            found += chosen(
                explain(told, portal, prior, shared, told_negated), item, told, "value", multiple
            )
    return found


def searched_text(char: Characteristic, value: str) -> str:
    """«Да» ищем по названию: «из натуральной кожи» значит «Натуральная кожа: Да». Слова
    шаблона названий вроде «Наличие» люди не пишут, их пропуск не штрафуем."""
    if value in ("Да", "Нет"):
        return " ".join(w["text"] for w in words(char["name"]) if w["text"] not in template())
    return value


@cache
def template() -> set[str]:
    """Слова, которые стоят больше чем в половине названий характеристик «Да/Нет» КТРУ."""
    names = {
        char["name"]
        for position in reference()
        for char in position["characteristics"]
        if char["values"] and all(v in ("Да", "Нет") for v in char["values"])
    }
    seen = Counter(text for name in names for text in {w["text"] for w in words(name)})
    return {text for text, count in seen.items() if count > len(names) / 2}


def number_values(
    fragments: list[Fragment],
    chars: list[Characteristic],
    portal_chars: list[Characteristic],
    allowed: dict[str, list[Value]],
    product: set[str],
) -> list[FoundChar]:
    """Числа в единице КТРУ: «2,2 см» → Толщина столешницы 22 мм, «высота 75» → 75 см."""
    numeric = numeric_chars(chars)
    portal_numeric = numeric_chars(portal_chars)
    found: list[FoundChar] = []
    for item in fragments:
        if item["form"] != "number" or item["number"] is None:
            continue
        name_stems = stems(item["name"] or "")
        char = owner(numeric, item["unit"], name_stems, guess=True)
        source: Source = "ktru"
        if char is None:
            listed = listed_number(chars, item, name_stems)
            if listed:
                found.append(listed)
                continue
        if char is None:
            # В названиях портала много мусора, поэтому название должно быть в тексте целиком.
            told = name_stems - product
            named_here = [c for c in portal_numeric if stems(c["name"]) <= told]
            char = owner(named_here, item["unit"], told, guess=False)
            source = "portal"
        if char is None:
            continue
        values = allowed.get(char["name"], [])
        char_found = number_char(
            char, source, "number", item["number"], item["unit"], values, item["text"]
        )
        char_found["start"], char_found["end"] = item["start"], item["end"]
        found.append(char_found)
    return found


def listed_number(
    chars: list[Characteristic], item: Fragment, name_stems: set[str]
) -> FoundChar | None:
    """«размер 48» → Российский размер: 48. В КТРУ размер это текст, но значение в списке."""
    said = show(tidy(item["number"] or 0))
    for char in chars:
        if char["type"] == "text" and said in char["values"] and stems(char["name"]) & name_stems:
            char_found = found_char(char, said, "ktru", "number", item["text"])
            char_found["start"], char_found["end"] = item["start"], item["end"]
            return char_found
    return None


def numeric_chars(chars: list[Characteristic]) -> list[Characteristic]:
    numeric = {}
    for char in chars:
        if char["type"] == "number":
            numeric[char["name"]] = char if "unit" in char else named_unit(char)
    return list(numeric.values())


def named_unit(char: Characteristic) -> Characteristic:
    """В КТРУ единица бывает только в названии: «Максимальный крутящий момент (Нм)»."""
    found = UNIT_IN_NAME.search(char["name"])
    if found is None or found["unit"].lower() not in UNIT_ALIASES:
        return char
    with_unit = char.copy()
    with_unit["unit"] = {"code": UNIT_ALIASES[found["unit"].lower()], "name": found["unit"]}
    return with_unit


def number_char(
    char: Characteristic,
    source: Source,
    rule: Rule,
    number: float,
    code: str | None,
    allowed: list[Value],
    fragment: str,
) -> FoundChar:
    unit: Unit | None = char["unit"] if "unit" in char else None
    value = number
    if unit and code:
        value = convert(number, code, unit["code"])
    item = found_char(char, tidy(value), source, rule, fragment, unit=unit)
    if unit and code and code != unit["code"]:
        mark(item, "unit")
    check_range(item, allowed)
    return item


def check_range(item: FoundChar, allowed: list[Value]) -> None:
    """Число вне коридора КТРУ: правим, если сдвиг на порядок его объясняет, иначе помечаем."""
    corridor = bounds(allowed)
    value = item["value"]
    if corridor is None or isinstance(value, str) or within(value, corridor):
        return
    for shift in SHIFTS:
        if within(value * shift, corridor):
            item["value"] = tidy(value * shift)
            mark(item, "wrong_unit")
            return
    low, high = corridor
    mark(item, "out_of_range", low=show(low), high=show(high))


def within(value: float, corridor: tuple[float, float]) -> bool:
    return corridor[0] <= value <= corridor[1]


def owner(
    chars: list[Characteristic], code: str | None, name_stems: set[str], guess: bool
) -> Characteristic | None:
    """Характеристика той же величины, чьё название ближе к словам перед числом. Угадывать
    без подсказки в тексте можно, только если выбора нет, единица названа и guess разрешён."""
    best = None
    best_hits = 0
    same_dimension = 0
    for char in chars:
        if code is not None and (
            "unit" not in char or scale(char["unit"]["code"])[0] != scale(code)[0]
        ):
            continue
        same_dimension += 1
        hits = len(stems(char["name"]) & name_stems)
        if best is None or hits > best_hits:
            best = char
            best_hits = hits
    if best_hits == 0 and (not guess or same_dimension > 1 or code is None):
        return None
    return best
