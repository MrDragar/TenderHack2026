"""Раздел это argmax P(раздел | n-граммы названия): наивный Байес на слабых метках листа 1.

Метка товара листа 1 это единственный раздел, у которого есть все названия его характеристик.
Учимся на товарах вне выборки замера и на самих названиях разделов. Байес считает каждую
букву столько раз, во сколько n-грамм она входит, и потому самоуверен: делим логарифм
правдоподобия на температуру T, подобранную на отложенной десятой части. Отвечаем, только
если P ≥ 1/2: неверный ответ стоит вдвое дороже отказа, порог следует из цен.
"""

import json
import zlib
from functools import cache

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar
from scipy.special import softmax
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics import log_loss
from sklearn.model_selection import GroupShuffleSplit
from sklearn.naive_bayes import MultinomialNB
from sklearn.pipeline import Pipeline, make_pipeline

from contracts.ref import Portal
from pipeline.data import portal as whole_portal
from scripts.products import (
    PRODUCTS_PATH,
    RAW_DIR,
    SAMPLE,
    SHEETS,
    section_names,
    split_pairs,
    weak_section,
)

TRAIN_PATH = PRODUCTS_PATH.with_name("section_train.json")
MIN_PROBA = 0.5
OTHER = ""


def find_section(text: str, portal: Portal) -> str | None:
    if not portal or not text.strip():
        return None
    model, temperature = trained(tuple(sorted(portal)))
    proba = softmax(model.predict_log_proba([text])[0] / temperature)
    best = proba.argmax()
    if proba[best] < MIN_PROBA or model.classes_[best] == OTHER:
        return None
    return str(model.classes_[best])


@cache
def trained(sections: tuple[str, ...]) -> tuple[Pipeline, float]:
    found, sorts = labeled()
    names = np.array(found)
    # Товары чужих разделов это исход «не из этих разделов»: без него модели некуда отказаться.
    labels = np.array([sort if sort in sections else OTHER for sort in sorts])
    # Одинаковые названия по одну сторону, иначе температура меряется на выученном.
    split = GroupShuffleSplit(n_splits=1, test_size=0.1, random_state=0)
    fit, held = next(split.split(names, labels, groups=names))
    # Название раздела весит как все товары среднего раздела: справочник и поставщики на равных.
    weights = [1.0] * len(fit) + [len(fit) / len(sections)] * len(sections)
    model = make_pipeline(
        TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), min_df=5, sublinear_tf=True),
        MultinomialNB(alpha=0.01),
    )
    model.fit(
        [*names[fit], *sections],
        [*labels[fit], *sections],
        multinomialnb__sample_weight=weights,
    )
    known = held[np.isin(labels[held], model.classes_)]
    joint = model.predict_log_proba(names[known])

    def loss(temperature: float) -> float:
        return log_loss(labels[known], softmax(joint / temperature, axis=1), labels=model.classes_)

    return model, minimize_scalar(loss, bounds=(0.1, 50), method="bounded").x


def labeled() -> tuple[list[str], list[str]]:
    if not TRAIN_PATH.exists():
        TRAIN_PATH.write_text(json.dumps(read_sheets(set(whole_portal())), ensure_ascii=False))
    rows = json.loads(TRAIN_PATH.read_text())
    return [name for name, _ in rows], [section for _, section in rows]


def read_sheets(portal: set[str]) -> list[tuple[str, str]]:
    """Товары, которые не попали в выборку замера: иначе учимся на том, по чему меряем."""
    rows = []
    for sheet_name in SHEETS:
        book = load_workbook(RAW_DIR / sheet_name, read_only=True)
        goods, chars = book.worksheets
        sections = section_names(
            chars.iter_rows(min_row=2, values_only=True), dict.fromkeys(portal)
        )
        names = set().union(*sections.values())
        for row in goods.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1] or not zlib.crc32(str(row[0]).encode()) % SAMPLE:
                continue
            section = weak_section(split_pairs(str(row[5] or ""), names), sections)
            if section:
                rows.append((" ".join(str(row[1]).split()), section))
        book.close()
    return rows
